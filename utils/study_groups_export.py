"""PDF, Excel and HD-image exports for a Study Groups set. The image export
(``build_png``) is drawn server-side with Pillow -- a modern indigo/violet
gradient hero banner, chip-based summary stats and a grid of group cards with
pastel avatar chips -- it is not a screenshot of the rendered page."""
import io

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from utils.numfmt import fmt_num as _n


def build_pdf(data, school):
    """A printable sheet: one table per group, leader marked, for the given
    ``_set_payload()`` data dict."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=12 * mm, bottomMargin=12 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm)
    base = getSampleStyleSheet()
    title_style = ParagraphStyle('t', parent=base['Title'], fontSize=16, alignment=TA_CENTER)
    sub_style = ParagraphStyle('s', parent=base['Normal'], fontSize=10, alignment=TA_CENTER,
                               textColor=colors.HexColor('#444444'))
    group_hdr = ParagraphStyle('g', parent=base['Normal'], fontSize=12, fontName='Helvetica-Bold',
                               spaceBefore=10, spaceAfter=4)

    flow = []
    school_name = (school or {}).get('name') or 'School'
    flow.append(Paragraph(school_name, title_style))
    flow.append(Paragraph(data.get('title') or 'Study Groups', sub_style))
    flow.append(Paragraph(
        f"{data.get('class_name', '')} &middot; {data.get('term_name', '')} "
        f"&middot; {data.get('num_groups')} group(s) of {data.get('group_size')}", sub_style))
    flow.append(Spacer(1, 6 * mm))

    for g in data.get('groups', []):
        flow.append(Paragraph(g['label'], group_hdr))
        rows = [['#', 'Student', 'Admission No.', 'Role']]
        for i, m in enumerate(g['members'], start=1):
            role = 'Leader' if m['is_leader'] else ''
            rows.append([str(i), m['name'], m['admission_no'], role])
        t = Table(rows, colWidths=[10 * mm, 80 * mm, 40 * mm, 25 * mm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        flow.append(t)
        flow.append(Spacer(1, 4 * mm))

    doc.build(flow)
    buf.seek(0)
    return buf


def build_xlsx(data):
    """One row per student: group, leader flag, basis average."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = 'Study Groups'
    head_font = Font(bold=True, color='FFFFFF')
    ws.append([data.get('title') or 'Study Groups'])
    ws.append([f"{data.get('class_name', '')} - {data.get('term_name', '')}"])
    ws.append([])
    header = ['Group', 'Student', 'Admission No.', 'Leader', 'Previous Term Average']
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center')
    for g in data.get('groups', []):
        for m in g['members']:
            ws.append([g['label'], m['name'], m['admission_no'],
                      'Yes' if m['is_leader'] else '',
                      _n(m['basis_average']) if m['basis_average'] is not None else ''])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_png(data, school):
    """A branded HD PNG of the board drawn server-side with Pillow: an indigo/
    violet gradient hero banner (logo + school identity), a chip-based summary
    with custom stat icons, and a grid of group cards -- each with a distinct
    gradient accent strip, pastel per-student avatar chips, a gold leader
    badge, and emerald/sky/rose/slate score pills. Not a page screenshot."""
    import os
    import math
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    from utils.student_export import _FONT_REG, _FONT_BOLD, _FONT_ITAL
    from utils.numfmt import fmt_num as n
    from utils import timeutil

    # ---- palette: indigo/violet brand, amber reserved for the "leader" signal ----
    INK = (15, 23, 42)
    SUBTLE = (71, 85, 105)
    MUTED = (100, 116, 139)
    LINE = (228, 232, 240)
    LINE_SOFT = (240, 242, 247)
    WHITE = (255, 255, 255)
    PAGE_BG = (250, 250, 252)

    INDIGO_DARK = (67, 56, 202)
    VIOLET = (124, 58, 237)
    INDIGO_BG = (238, 237, 253)
    INDIGO_FG = (67, 56, 202)

    AMBER = (245, 158, 11)
    AMBER_WASH = (255, 251, 235)
    AMBER_CHIP_BG = (254, 243, 199)
    AMBER_CHIP_FG = (146, 64, 14)

    EMERALD_BG, EMERALD_FG = (209, 250, 229), (4, 120, 87)
    SKY_BG, SKY_FG = (224, 242, 254), (3, 105, 161)
    ROSE_BG, ROSE_FG = (255, 228, 230), (159, 18, 57)
    SLATE_BG, SLATE_FG = (241, 245, 249), (100, 116, 139)

    AVATAR_PALETTE = [
        ((224, 231, 255), (67, 56, 202)),   # indigo
        ((255, 228, 230), (190, 18, 60)),   # rose
        ((220, 252, 231), (21, 128, 61)),   # green
        ((254, 249, 195), (161, 98, 7)),    # yellow
        ((224, 242, 254), (3, 105, 161)),   # sky
        ((243, 232, 255), (107, 33, 168)),  # purple
        ((255, 237, 213), (194, 65, 12)),   # orange
        ((204, 251, 241), (15, 118, 110)),  # teal
    ]
    CARD_GRADIENTS = [
        ((79, 70, 229), (124, 58, 237)),    # indigo -> violet
        ((14, 165, 233), (79, 70, 229)),    # sky -> indigo
        ((124, 58, 237), (219, 39, 119)),   # violet -> pink
        ((16, 185, 129), (14, 165, 233)),   # emerald -> sky
        ((20, 184, 166), (79, 70, 229)),    # teal -> indigo
    ]

    S = 2                       # supersample factor; downscaled at the end for anti-aliasing
    BASE_W = 1480
    W = BASE_W * S
    margin = 44 * S

    def fnt(size, bold=False, italic=False):
        p = _FONT_BOLD if bold else (_FONT_ITAL if italic else _FONT_REG)
        try:
            return ImageFont.truetype(p, int(size * S))
        except Exception:
            return ImageFont.load_default()

    probe = ImageDraw.Draw(Image.new('RGB', (1, 1)))

    def tw(t, f):
        b = probe.textbbox((0, 0), str(t), font=f)
        return b[2] - b[0]

    def th(f):
        b = probe.textbbox((0, 0), 'Ay', font=f)
        return b[3] - b[1]

    def fit(t, f, mw):
        t = str(t)
        if tw(t, f) <= mw:
            return t
        while t and tw(t + '…', f) > mw:
            t = t[:-1]
        return (t + '…') if t else ''

    def initials(name):
        parts = [p for p in str(name).split() if p]
        if not parts:
            return '?'
        if len(parts) == 1:
            return parts[0][:2].upper()
        return (parts[0][0] + parts[-1][0]).upper()

    def score_tier(avg):
        if avg is None:
            return SLATE_BG, SLATE_FG, 'NEW'
        v = float(avg)
        txt = n(avg)
        if v >= 70:
            return EMERALD_BG, EMERALD_FG, txt
        if v >= 50:
            return SKY_BG, SKY_FG, txt
        return ROSE_BG, ROSE_FG, txt

    def star_points(cx, cy, r_out, r_in, rot=-90):
        pts = []
        for i in range(10):
            ang = math.radians(rot + i * 36)
            r = r_out if i % 2 == 0 else r_in
            pts.append((cx + r * math.cos(ang), cy + r * math.sin(ang)))
        return pts

    def spaced_w(text, f, tracking):
        if not text:
            return 0
        return sum(tw(c, f) for c in text) + tracking * (len(text) - 1)

    def draw_spaced(dctx, xy, text, f, fill, tracking):
        x, y = xy
        for c in text:
            dctx.text((x, y), c, font=f, fill=fill)
            x += tw(c, f) + tracking

    def diagonal_gradient(w, h, c1, c2):
        """Fast smooth diagonal gradient via a tiny grid upscaled with bilinear filtering."""
        w, h = max(1, int(w)), max(1, int(h))
        grid = 40
        small = Image.new('RGB', (grid, grid))
        px = small.load()
        for yy in range(grid):
            for xx in range(grid):
                t = (xx + yy) / (2 * (grid - 1))
                px[xx, yy] = tuple(int(c1[k] + (c2[k] - c1[k]) * t) for k in range(3))
        return small.resize((w, h), Image.BILINEAR)

    def rounded_mask(w, h, radius, corners=(True, True, True, True)):
        w, h = max(1, int(w)), max(1, int(h))
        mask = Image.new('L', (w, h), 0)
        md = ImageDraw.Draw(mask)
        md.rounded_rectangle([0, 0, w - 1, h - 1], radius=radius, fill=255, corners=corners)
        return mask

    def shadow(img, box, radius, blur=16, alpha=26, offset=(0, 8 * S)):
        x0, y0, x1, y1 = box
        pad = blur * 3
        sw, sh = int(x1 - x0 + 2 * pad), int(y1 - y0 + 2 * pad)
        if sw <= 0 or sh <= 0:
            return
        sh_img = Image.new('RGBA', (sw, sh), (0, 0, 0, 0))
        sd = ImageDraw.Draw(sh_img)
        sd.rounded_rectangle([pad, pad, sw - pad, sh - pad], radius=radius, fill=(30, 32, 58, alpha))
        sh_img = sh_img.filter(ImageFilter.GaussianBlur(blur))
        img.paste(sh_img, (int(x0 - pad + offset[0]), int(y0 - pad + offset[1])), sh_img)

    def shadow_ellipse(img, cx, cy, r, blur=4, alpha=28, offset=(0, 2 * S)):
        pad = blur * 3
        size = int(r * 2 + 2 * pad)
        sh_img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
        sd = ImageDraw.Draw(sh_img)
        sd.ellipse([pad, pad, size - pad, size - pad], fill=(20, 22, 45, alpha))
        sh_img = sh_img.filter(ImageFilter.GaussianBlur(blur))
        img.paste(sh_img, (int(cx - r - pad + offset[0]), int(cy - r - pad + offset[1])), sh_img)

    def icon_groups(dctx, cx, cy, s, color):
        g, sq = s * 0.22, s * 0.38
        for ox in (-1, 1):
            for oy in (-1, 1):
                x0 = cx + ox * (g / 2 + sq / 2) - sq / 2
                y0 = cy + oy * (g / 2 + sq / 2) - sq / 2
                dctx.rounded_rectangle([x0, y0, x0 + sq, y0 + sq], radius=sq * 0.26, fill=color)

    def icon_people(dctx, cx, cy, s, color, bg):
        r = s * 0.22
        for ox, fill in ((-s * 0.16, bg), (s * 0.16, color)):
            hx, hy = cx + ox, cy - s * 0.12
            dctx.ellipse([hx - r, hy - r, hx + r, hy + r], fill=fill, outline=color, width=max(1, int(1 * S)))
            bw, bh = s * 0.58, s * 0.3
            dctx.rounded_rectangle([hx - bw / 2, hy + r - 1, hx + bw / 2, hy + r + bh], radius=bh * 0.4,
                                   fill=fill, outline=color, width=max(1, int(1 * S)))

    def icon_calendar(dctx, cx, cy, s, color):
        w, h = s * 0.62, s * 0.56
        x0, y0 = cx - w / 2, cy - h / 2 + s * 0.04
        dctx.rounded_rectangle([x0, y0, x0 + w, y0 + h], radius=s * 0.08, outline=color, width=max(1, int(2 * S)))
        dctx.line([x0, y0 + h * 0.34, x0 + w, y0 + h * 0.34], fill=color, width=max(1, int(2 * S)))
        for fx in (x0 + w * 0.26, x0 + w * 0.74):
            dctx.line([fx, y0 - s * 0.06, fx, y0 + h * 0.18], fill=color, width=max(1, int(2 * S)))

    # Fonts
    name_f = fnt(30, True)
    addr_f = fnt(13)
    title_f = fnt(28, True)
    meta_f = fnt(13)
    attrib_f = fnt(11)
    chip_f = fnt(11, True)
    stat_lab_f, stat_val_f = fnt(10, True), fnt(17, True)
    hdr_f = fnt(16, True)
    count_f = fnt(11, True)
    member_f, member_b = fnt(13), fnt(13, True)
    avatar_f = fnt(12, True)
    badge_f = fnt(10, True)
    score_f = fnt(11, True)
    foot_s = fnt(9)

    groups = data.get('groups') or []
    title = data.get('title') or 'Study Groups'
    basis_term_name = data.get('basis_term_name')
    meta_line = (f"{data.get('class_name', '')}   ·   {data.get('term_name', '')}   ·   "
                f"{data.get('num_groups')} group(s) of {data.get('group_size')}")
    rank_line = (f"Ranked by {basis_term_name} average" if basis_term_name
                else 'No prior-term data available — new students placed randomly')
    attrib_line = f"Generated by {data.get('created_by', '')} on {data.get('created_at', '')}"

    # ---- sizing passes ----
    banner_h = 148 * S
    sum_top = banner_h + 28 * S
    sum_h = 168 * S
    sum_pad = 32 * S

    gap = 22 * S
    min_card_w = 340 * S
    avail = W - 2 * margin
    cols = max(1, min(4, int((avail + gap) // (min_card_w + gap))))
    card_w = int((avail - (cols - 1) * gap) / cols)

    strip_h = 7 * S
    hdr_pad_top = 20 * S
    hdr_row_h = 44 * S
    member_row_h = int(th(member_f) * 3.0)
    card_radius = 20 * S
    avatar_d = 34 * S
    pad_x = 20 * S

    def card_height(g):
        return strip_h + hdr_pad_top + hdr_row_h + max(1, len(g.get('members') or [])) * member_row_h + 10 * S

    rows_layout = [groups[i:i + cols] for i in range(0, len(groups), cols)] or [[]]
    grid_top = sum_top + sum_h + 32 * S
    y = grid_top
    row_tops = []
    for chunk in rows_layout:
        h = max((card_height(g) for g in chunk), default=0)
        row_tops.append((y, h))
        y += h + gap
    grid_bottom = (y - gap) if groups else grid_top
    foot_h = 54 * S
    H = int(grid_bottom + foot_h + margin)

    img = Image.new('RGB', (W, H), PAGE_BG)
    d = ImageDraw.Draw(img)

    # ---- hero banner: indigo -> violet diagonal gradient with soft glow blobs ----
    banner = diagonal_gradient(W, banner_h, INDIGO_DARK, VIOLET)
    glow = Image.new('RGBA', (W, int(banner_h)), (0, 0, 0, 0))
    gd = ImageDraw.Draw(glow)
    gd.ellipse([W * 0.72, -banner_h * 0.6, W * 0.72 + banner_h * 2.0, banner_h * 1.4],
              fill=(255, 255, 255, 22))
    gd.ellipse([W * 0.92, banner_h * 0.1, W * 0.92 + banner_h * 1.1, banner_h * 1.3],
              fill=(255, 255, 255, 16))
    glow = glow.filter(ImageFilter.GaussianBlur(40))
    banner = Image.alpha_composite(banner.convert('RGBA'), glow).convert('RGB')
    img.paste(banner, (0, 0))

    bx = margin
    logo = (school or {}).get('logo_path')
    circ_d = 96 * S
    circ_y = (banner_h - circ_d) / 2
    logo_drawn = False
    if logo and os.path.exists(logo):
        try:
            d.ellipse([bx, circ_y, bx + circ_d, circ_y + circ_d], fill=WHITE)
            lg = Image.open(logo).convert('RGBA')
            inner = circ_d - 16 * S
            w0, h0 = lg.size
            scale = inner / max(w0, h0)
            lg = lg.resize((max(1, int(w0 * scale)), max(1, int(h0 * scale))), Image.LANCZOS)
            lx = bx + (circ_d - lg.width) / 2
            ly = circ_y + (circ_d - lg.height) / 2
            img.paste(lg, (int(lx), int(ly)), lg)
            logo_drawn = True
        except Exception:
            logo_drawn = False
    if not logo_drawn:
        d.ellipse([bx, circ_y, bx + circ_d, circ_y + circ_d], fill=WHITE)
        mono = ''.join(w[0] for w in ((school or {}).get('name') or 'S').split()[:2]).upper()
        mf = fnt(26, True)
        d.text((bx + circ_d / 2 - tw(mono, mf) / 2, circ_y + circ_d / 2 - th(mf) / 2), mono,
               fill=INDIGO_DARK, font=mf)
    text_x = bx + circ_d + 26 * S

    name_right_limit = W - margin
    nm = ((school or {}).get('name') or 'School').upper()
    nf, nsz = name_f, 30
    while nsz > 18 and tw(nm, nf) > (name_right_limit - text_x):
        nsz -= 2
        nf = fnt(nsz, True)
    name_top = banner_h / 2 - 42 * S
    d.text((text_x, name_top), fit(nm, nf, name_right_limit - text_x), fill=WHITE, font=nf)
    ty = name_top + th(nf) + 14 * S
    addr = (school or {}).get('address') or ''
    contact = '      '.join(p for p in [(school or {}).get('phone') or '', (school or {}).get('email') or ''] if p)
    line2 = '   ·   '.join(p for p in [addr, contact] if p)
    if line2:
        d.text((text_x, ty), fit(line2, addr_f, name_right_limit - text_x), fill=(221, 216, 254), font=addr_f)
        ty += th(addr_f) + 10 * S
    motto = (school or {}).get('motto') or ''
    if motto:
        d.text((text_x, ty), fit(motto, addr_f, name_right_limit - text_x), fill=(197, 190, 248), font=addr_f)

    # ---- summary: title + report chip (left), stat chips (right) -- flat, chip-based ----
    shadow(img, (margin, sum_top, W - margin, sum_top + sum_h), card_radius)
    d.rounded_rectangle([margin, sum_top, W - margin, sum_top + sum_h], radius=card_radius, fill=WHITE,
                        outline=LINE, width=1)

    stat_w, stat_h, stat_gap = 148 * S, sum_h - 48 * S, 14 * S
    n_stats = 3
    stats_total_w = stat_w * n_stats + stat_gap * (n_stats - 1)
    stats_x0 = W - margin - sum_pad - stats_total_w
    text_right_limit = stats_x0 - 36 * S

    tx0 = margin + sum_pad
    ty0 = sum_top + 26 * S
    tt = fit(str(title), title_f, text_right_limit - tx0)
    d.text((tx0, ty0), tt, fill=INK, font=title_f)
    chip_txt = 'GROUP ASSIGNMENT REPORT'
    ctrack = 1.2 * S
    cw = spaced_w(chip_txt, chip_f, ctrack) + 24 * S
    ch_ = 24 * S
    cx0 = tx0 + tw(tt, title_f) + 16 * S
    cy0 = ty0 + th(title_f) / 2 - ch_ / 2
    if cx0 + cw <= text_right_limit:
        d.rounded_rectangle([cx0, cy0, cx0 + cw, cy0 + ch_], radius=ch_ / 2, fill=INDIGO_BG)
        draw_spaced(d, (cx0 + 12 * S, cy0 + ch_ / 2 - th(chip_f) / 2), chip_txt, chip_f, INDIGO_FG, ctrack)
    ty0 += th(title_f) + 16 * S
    d.text((tx0, ty0), fit(meta_line, meta_f, text_right_limit - tx0), fill=SUBTLE, font=meta_f)
    ty0 += th(meta_f) + 10 * S
    d.text((tx0, ty0), fit(rank_line, meta_f, text_right_limit - tx0), fill=MUTED, font=meta_f)
    ty0 += th(meta_f) + 12 * S
    d.text((tx0, ty0), fit(attrib_line, attrib_f, text_right_limit - tx0), fill=MUTED, font=attrib_f)

    stat_cells = [('Groups', str(data.get('num_groups') or len(groups)), 'g'),
                 ('Per group', str(data.get('group_size') or ''), 'p'),
                 ('Date', timeutil.today().strftime('%d %b %Y'), 'd')]
    sy = sum_top + 24 * S
    for i, (lab, val, kind) in enumerate(stat_cells):
        sx = stats_x0 + i * (stat_w + stat_gap)
        d.rounded_rectangle([sx, sy, sx + stat_w, sy + stat_h], radius=16 * S, fill=INDIGO_BG)
        icon_cx, icon_cy = sx + 30 * S, sy + stat_h / 2
        if kind == 'g':
            icon_groups(d, icon_cx, icon_cy, 24 * S, INDIGO_FG)
        elif kind == 'p':
            icon_people(d, icon_cx, icon_cy, 24 * S, INDIGO_FG, INDIGO_BG)
        else:
            icon_calendar(d, icon_cx, icon_cy, 24 * S, INDIGO_FG)
        tx = sx + 54 * S
        d.text((tx, sy + stat_h / 2 - th(stat_val_f) - 2 * S), val, fill=INK, font=stat_val_f)
        d.text((tx, sy + stat_h / 2 + 4 * S), lab, fill=MUTED, font=stat_lab_f)

    # ---- group cards ----
    for row_i, chunk in enumerate(rows_layout):
        row_y, row_h = row_tops[row_i] if groups else (grid_top, 0)
        for col_i, g in enumerate(chunk):
            card_idx = row_i * cols + col_i
            cx0 = margin + col_i * (card_w + gap)
            ch = card_height(g)
            shadow(img, (cx0, row_y, cx0 + card_w, row_y + ch), card_radius, blur=14, alpha=22)
            d.rounded_rectangle([cx0, row_y, cx0 + card_w, row_y + ch], radius=card_radius, fill=WHITE,
                                outline=LINE, width=1)
            g1, g2 = CARD_GRADIENTS[card_idx % len(CARD_GRADIENTS)]
            strip = diagonal_gradient(card_w, strip_h, g1, g2)
            mask = rounded_mask(card_w, strip_h * 3, card_radius, corners=(True, True, False, False))
            mask = mask.crop((0, 0, card_w, strip_h))
            img.paste(strip, (int(cx0), int(row_y)), mask)

            members = g.get('members') or []
            label_y = row_y + strip_h + hdr_pad_top
            label = fit(g.get('label', ''), hdr_f, card_w - 2 * pad_x - 60 * S)
            d.text((cx0 + pad_x, label_y), label, fill=INK, font=hdr_f)
            cnt_txt = f"{len(members)}"
            pill_w = max(40 * S, tw(cnt_txt, count_f) + 24 * S)
            pill_h = 26 * S
            pill_x = cx0 + card_w - pad_x - pill_w
            pill_y = label_y + th(hdr_f) / 2 - pill_h / 2
            d.rounded_rectangle([pill_x, pill_y, pill_x + pill_w, pill_y + pill_h], radius=pill_h / 2,
                                fill=INDIGO_BG)
            d.text((pill_x + pill_w / 2 - tw(cnt_txt, count_f) / 2, pill_y + pill_h / 2 - th(count_f) / 2),
                   cnt_txt, fill=INDIGO_FG, font=count_f)

            yy = row_y + strip_h + hdr_pad_top + hdr_row_h
            if not members:
                d.text((cx0 + pad_x, yy + 14 * S), 'No members', fill=MUTED, font=member_f)
            for idx, m in enumerate(members):
                is_leader = bool(m.get('is_leader'))
                if is_leader:
                    d.rectangle([cx0 + 2, yy, cx0 + card_w - 2, yy + member_row_h], fill=AMBER_WASH)
                    d.rectangle([cx0 + 2, yy, cx0 + 4 * S, yy + member_row_h], fill=AMBER)

                av_cx = cx0 + pad_x + avatar_d / 2 + (4 * S if is_leader else 0)
                av_cy = yy + member_row_h / 2
                shadow_ellipse(img, av_cx, av_cy, avatar_d / 2)
                if is_leader:
                    av_fill, av_fg = AMBER, WHITE
                else:
                    av_fill, av_fg = AVATAR_PALETTE[idx % len(AVATAR_PALETTE)]
                d.ellipse([av_cx - avatar_d / 2, av_cy - avatar_d / 2, av_cx + avatar_d / 2, av_cy + avatar_d / 2],
                          fill=av_fill, outline=WHITE, width=max(1, int(2 * S)))
                ini = initials(m.get('name', ''))
                d.text((av_cx - tw(ini, avatar_f) / 2, av_cy - th(avatar_f) / 2), ini, fill=av_fg, font=avatar_f)
                if is_leader:
                    bd_r = 10 * S
                    bd_cx, bd_cy = av_cx + avatar_d / 2 - 2 * S, av_cy - avatar_d / 2 + 2 * S
                    d.ellipse([bd_cx - bd_r, bd_cy - bd_r, bd_cx + bd_r, bd_cy + bd_r], fill=WHITE,
                             outline=AMBER, width=2)
                    pts = star_points(bd_cx, bd_cy, bd_r * 0.62, bd_r * 0.26)
                    d.polygon(pts, fill=AMBER)

                name_x = av_cx + avatar_d / 2 + 14 * S
                bg_c, fg_c, score_txt = score_tier(m.get('basis_average'))
                sp_w = max(48 * S, tw(score_txt, score_f) + 22 * S)
                sp_h = 25 * S
                sp_x = cx0 + card_w - pad_x - sp_w
                sp_y = yy + (member_row_h - sp_h) / 2
                d.rounded_rectangle([sp_x, sp_y, sp_x + sp_w, sp_y + sp_h], radius=sp_h / 2, fill=bg_c)
                d.text((sp_x + sp_w / 2 - tw(score_txt, score_f) / 2, sp_y + sp_h / 2 - th(score_f) / 2),
                       score_txt, fill=fg_c, font=score_f)

                badge_w = 0
                if is_leader:
                    badge_txt = 'LEADER'
                    badge_w = tw(badge_txt, badge_f) + 18 * S
                name_mw = sp_x - 10 * S - name_x - (badge_w + 8 * S if badge_w else 0)
                name_font = member_b if is_leader else member_f
                name_txt = fit(m.get('name', ''), name_font, max(10 * S, name_mw))
                tyy = yy + (member_row_h - th(name_font)) / 2
                d.text((name_x, tyy), name_txt, fill=INK, font=name_font)
                if badge_w:
                    bxp = name_x + tw(name_txt, name_font) + 8 * S
                    bh2 = 20 * S
                    byp = yy + (member_row_h - bh2) / 2
                    d.rounded_rectangle([bxp, byp, bxp + badge_w, byp + bh2], radius=bh2 / 2, fill=AMBER_CHIP_BG)
                    d.text((bxp + badge_w / 2 - tw('LEADER', badge_f) / 2, byp + bh2 / 2 - th(badge_f) / 2),
                           'LEADER', fill=AMBER_CHIP_FG, font=badge_f)

                if idx < len(members) - 1:
                    d.line([cx0 + pad_x, yy + member_row_h, cx0 + card_w - pad_x, yy + member_row_h],
                           fill=LINE_SOFT, width=1)
                yy += member_row_h

    # ---- footer ----
    fy = H - margin - foot_h + 14 * S
    d.line([margin, fy, W - margin, fy], fill=LINE, width=1)
    school_name = (school or {}).get('name') or 'School'
    d.text((margin, fy + 14 * S), school_name, fill=SUBTLE, font=foot_s)
    total_students = sum(len(g.get('members') or []) for g in groups)
    gist = f"{total_students} student(s) across {len(groups)} group(s)"
    d.text((W / 2 - tw(gist, foot_s) / 2, fy + 14 * S), gist, fill=MUTED, font=foot_s)
    conf = 'Confidential — for school, parent and guardian use only'
    d.text((W - margin - tw(conf, foot_s), fy + 14 * S), conf, fill=MUTED, font=foot_s)

    out_w, out_h = BASE_W, max(1, H // S)
    buf = io.BytesIO()
    img.resize((out_w, out_h), Image.LANCZOS).save(buf, format='PNG')
    buf.seek(0)
    return buf
