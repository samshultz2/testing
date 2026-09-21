"""Timetable generator: co-schedule rules pair two subjects (usually from
different arms of a combined class) into the exact same slot every time —
the mirror of the existing subject-clash rules, which forbid that."""
from config import Config
from models import (
    db, Branch, GenSubject, GenTeacher, GenTeacherAssignment, GenClassConfig,
    GenClassArmStream, GenStream, GenStreamSubject, GenSubjectConfig,
    GenCoScheduleRule, GenTimetableResult,
)
from tests.conftest import login_token


def _admin(app):
    c = app.test_client()
    c.post('/login', data={'password': Config.ADMIN_PASSWORD, '_csrf_token': login_token(c)})
    with c.session_transaction() as s:
        s['_csrf_token'] = 'a' * 64
    return c


def _post(c, url, **data):
    data.setdefault('_csrf_token', 'a' * 64)
    return c.post(url, data=data)


def _build_combined_class(app, tag):
    """SSS2 with two arms on two streams: Daisy=Arts (Literature),
    Iris=Commercial (Accounting) — the exact scenario from the feature
    request. Returns (class_config_id, literature_id, accounting_id).
    `tag` keeps names unique across tests sharing one session-scoped DB."""
    with app.app_context():
        bid = Branch.get_default().id

        cc = GenClassConfig(branch_id=bid, class_name=f'ZzSSS2{tag}', school_level='sss',
                            num_arms=2, arm_names=f'ZzDaisy{tag},ZzIris{tag}', has_streams=True)
        db.session.add(cc); db.session.flush()

        arts = GenStream(branch_id=bid, name=f'ZzArts{tag}', school_level='sss')
        comm = GenStream(branch_id=bid, name=f'ZzCommercial{tag}', school_level='sss')
        db.session.add_all([arts, comm]); db.session.flush()

        db.session.add_all([
            GenClassArmStream(class_config_id=cc.id, arm_name=f'ZzDaisy{tag}', stream_id=arts.id),
            GenClassArmStream(class_config_id=cc.id, arm_name=f'ZzIris{tag}', stream_id=comm.id),
        ])

        lit = GenSubject(branch_id=bid, name=f'ZzLiterature{tag}', school_level='sss')
        acct = GenSubject(branch_id=bid, name=f'ZzAccounting{tag}', school_level='sss')
        db.session.add_all([lit, acct]); db.session.flush()

        db.session.add_all([
            # Exempt from the (unrelated) day-separation default: this fixture is
            # about co-schedule pairing specifically, and coupling that with the
            # school-wide day-separation rule needlessly tightens an already
            # constrained (same-slot-paired) solve for no reason relevant to what
            # these tests check.
            GenSubjectConfig(branch_id=bid, subject_id=lit.id, school_level='sss', periods_per_week=2,
                             day_separation_exempt=True),
            GenSubjectConfig(branch_id=bid, subject_id=acct.id, school_level='sss', periods_per_week=2,
                             day_separation_exempt=True),
            GenStreamSubject(stream_id=arts.id, subject_id=lit.id, periods_per_week=2),
            GenStreamSubject(stream_id=comm.id, subject_id=acct.id, periods_per_week=2),
        ])

        lit_teacher = GenTeacher(branch_id=bid, name=f'Zz Lit Teacher{tag}', school_level='sss',
                                 max_periods_per_day=6, max_periods_per_week=30)
        acct_teacher = GenTeacher(branch_id=bid, name=f'Zz Acct Teacher{tag}', school_level='sss',
                                  max_periods_per_day=6, max_periods_per_week=30)
        db.session.add_all([lit_teacher, acct_teacher]); db.session.flush()

        db.session.add_all([
            GenTeacherAssignment(branch_id=bid, teacher_id=lit_teacher.id, subject_id=lit.id,
                                 class_config_id=cc.id, arm_name=f'ZzDaisy{tag}'),
            GenTeacherAssignment(branch_id=bid, teacher_id=acct_teacher.id, subject_id=acct.id,
                                 class_config_id=cc.id, arm_name=f'ZzIris{tag}'),
        ])
        db.session.commit()
        return cc.id, lit.id, acct.id


def test_coschedule_rule_model_and_routes(app):
    cc_id, lit_id, acct_id = _build_combined_class(app, 'A')
    c = _admin(app)

    r = _post(c, '/generator/coschedule-rules/add', name='ZzLit+Acct',
             source_subject_id=lit_id, source_class_name='ZzSSS2A', source_arm_name='ZzDaisyA',
             target_subject_id=acct_id, target_class_name='ZzSSS2A', target_arm_name='ZzIrisA')
    assert r.status_code in (302, 200)

    with app.app_context():
        rule = GenCoScheduleRule.query.filter_by(name='ZzLit+Acct').first()
        assert rule is not None
        assert rule.is_active
        assert rule.source_arm_name == 'ZzDaisyA' and rule.target_arm_name == 'ZzIrisA'
        rule_id = rule.id

    listing = c.get('/generator/clash-rules')
    assert listing.status_code == 200
    assert b'ZzLit+Acct' in listing.data

    r2 = _post(c, f'/generator/coschedule-rules/{rule_id}/toggle')
    assert r2.status_code in (302, 200)
    with app.app_context():
        assert GenCoScheduleRule.query.get(rule_id).is_active is False

    r3 = _post(c, f'/generator/coschedule-rules/{rule_id}/delete')
    assert r3.status_code in (302, 200)
    with app.app_context():
        assert GenCoScheduleRule.query.get(rule_id) is None


def test_coschedule_rule_requires_specific_arm(app):
    """Unlike clash rules, a co-schedule pairing can't use 'all arms' —
    pairing needs an exact 1:1 correspondence between two named groups."""
    cc_id, lit_id, acct_id = _build_combined_class(app, 'B')
    c = _admin(app)
    r = _post(c, '/generator/coschedule-rules/add', name='ZzBadRule',
             source_subject_id=lit_id, source_class_name='ZzSSS2B', source_arm_name='',
             target_subject_id=acct_id, target_class_name='ZzSSS2B', target_arm_name='ZzIrisB')
    assert r.status_code in (302, 200)
    with app.app_context():
        assert GenCoScheduleRule.query.filter_by(name='ZzBadRule').first() is None


def test_solver_pairs_co_scheduled_subjects_into_the_same_slot(app):
    """The actual feature request: force Literature (Daisy) and Accounting
    (Iris) into the same slot every time, with no teacher double-booking."""
    cc_id, lit_id, acct_id = _build_combined_class(app, 'C')
    c = _admin(app)

    with app.app_context():
        bid = Branch.get_default().id
        db.session.add(GenCoScheduleRule(
            branch_id=bid, name='ZzPair', source_subject_id=lit_id,
            source_class_name='ZzSSS2C', source_arm_name='ZzDaisyC',
            target_subject_id=acct_id, target_class_name='ZzSSS2C', target_arm_name='ZzIrisC',
            is_active=True))
        db.session.commit()

    r = c.post('/generator/generate/ortools',
              data={'_csrf_token': 'a' * 64, 'class_ids[]': cc_id, 'time_limit': '20', 'periods_per_day': '6'},
              follow_redirects=True)
    assert r.status_code == 200

    with app.app_context():
        # Scoped to this test's own arm names rather than "latest batch_id
        # overall" — the shared session-scoped DB means another test's batch
        # can otherwise look like the "latest" one and mask a real failure
        # here (or a real failure here can slip through by matching an
        # unrelated earlier batch's rows).
        rows = GenTimetableResult.query.filter(GenTimetableResult.arm_name.in_(['ZzDaisyC', 'ZzIrisC'])).all()
        assert rows, 'no timetable rows saved — generation likely failed; check flash message'
        batch_id = rows[0].batch_id
        batch_rows = [row for row in rows if row.batch_id == batch_id]

        lit_rows = [row for row in batch_rows if row.arm_name == 'ZzDaisyC' and row.subject_id == lit_id]
        acct_rows = [row for row in batch_rows if row.arm_name == 'ZzIrisC' and row.subject_id == acct_id]
        assert len(lit_rows) == 2 and len(acct_rows) == 2

        lit_slots = {(row.day_of_week, row.period_number) for row in lit_rows}
        acct_slots = {(row.day_of_week, row.period_number) for row in acct_rows}
        assert lit_slots == acct_slots, (
            f'Literature (Daisy) and Accounting (Iris) should land on identical slots, '
            f'got {lit_slots} vs {acct_slots}')

        # Each subject's own teacher isn't double-booked at that slot with
        # anything else (baseline solver guarantee — nothing coschedule-specific
        # needed for this, verified here as a sanity check on the pairing).
        for row in lit_rows + acct_rows:
            same_slot_same_teacher = [
                r2 for r2 in batch_rows
                if r2.day_of_week == row.day_of_week and r2.period_number == row.period_number
                and r2.teacher_id == row.teacher_id]
            assert len(same_slot_same_teacher) == 1


def test_print_results_can_filter_to_selected_arms(app):
    """The master 'Print All' view can be narrowed to specific class-arms —
    e.g. print only one side of a co-scheduled combined class instead of two
    near-identical tables."""
    cc_id, lit_id, acct_id = _build_combined_class(app, 'D')
    batch_id = 'zzbatch-print-filter'
    with app.app_context():
        bid = Branch.get_default().id
        db.session.add_all([
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name='ZzSSS2D', arm_name='ZzDaisyD', day_of_week=0,
                               period_number=1, subject_id=lit_id),
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name='ZzSSS2D', arm_name='ZzIrisD', day_of_week=0,
                               period_number=1, subject_id=acct_id),
        ])
        db.session.commit()

    c = _admin(app)

    # No filter — both arms show (existing "Print All" behaviour, unchanged).
    r_all = c.get(f'/generator/results/{batch_id}/print')
    assert r_all.status_code == 200
    body_all = r_all.get_data(as_text=True)
    assert 'ZzDaisyD' in body_all and 'ZzIrisD' in body_all

    # Filtered to just the Daisy arm.
    r_daisy = c.get(f'/generator/results/{batch_id}/print?arm=ZzSSS2D|ZzDaisyD')
    assert r_daisy.status_code == 200
    body_daisy = r_daisy.get_data(as_text=True)
    assert 'ZzDaisyD' in body_daisy and 'ZzIrisD' not in body_daisy

    # Filtered to just the Iris arm.
    r_iris = c.get(f'/generator/results/{batch_id}/print?arm=ZzSSS2D|ZzIrisD')
    assert r_iris.status_code == 200
    body_iris = r_iris.get_data(as_text=True)
    assert 'ZzIrisD' in body_iris and 'ZzDaisyD' not in body_iris


def _seed_print_filter_batch(app, tag):
    """A combined class's two arms, each with one result row, for exercising
    the shared _filter_by_arm() across print/export routes."""
    cc_id, lit_id, acct_id = _build_combined_class(app, tag)
    batch_id = f'zzbatch-export-filter-{tag}'
    with app.app_context():
        bid = Branch.get_default().id
        db.session.add_all([
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name=f'ZzSSS2{tag}', arm_name=f'ZzDaisy{tag}', day_of_week=0,
                               period_number=1, subject_id=lit_id),
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name=f'ZzSSS2{tag}', arm_name=f'ZzIris{tag}', day_of_week=0,
                               period_number=1, subject_id=acct_id),
        ])
        db.session.commit()
    return batch_id


def test_export_by_class_can_filter_to_selected_arms(app):
    import openpyxl
    from io import BytesIO

    batch_id = _seed_print_filter_batch(app, 'E')
    c = _admin(app)

    r_all = c.get(f'/generator/results/{batch_id}/export')
    assert r_all.status_code == 200
    wb_all = openpyxl.load_workbook(BytesIO(r_all.data))
    assert len(wb_all.sheetnames) == 2

    r_daisy = c.get(f'/generator/results/{batch_id}/export?arm=ZzSSS2E|ZzDaisyE')
    assert r_daisy.status_code == 200
    wb_daisy = openpyxl.load_workbook(BytesIO(r_daisy.data))
    assert len(wb_daisy.sheetnames) == 1
    assert 'Daisy' in wb_daisy.sheetnames[0] and 'Iris' not in wb_daisy.sheetnames[0]


def test_export_by_day_can_filter_to_selected_arms(app):
    import openpyxl
    from io import BytesIO

    batch_id = _seed_print_filter_batch(app, 'F')
    c = _admin(app)

    r_all = c.get(f'/generator/results/{batch_id}/export_by_day')
    assert r_all.status_code == 200
    wb_all = openpyxl.load_workbook(BytesIO(r_all.data))
    monday_all = wb_all['Monday']
    # Header rows (school name, address, day, periods) then one data row per arm.
    rows_with_class_code = sum(1 for row in monday_all.iter_rows() if row[0].value and row[0].value.strip())
    all_row_count = rows_with_class_code

    r_daisy = c.get(f'/generator/results/{batch_id}/export_by_day?arm=ZzSSS2F|ZzDaisyF')
    assert r_daisy.status_code == 200
    wb_daisy = openpyxl.load_workbook(BytesIO(r_daisy.data))
    monday_daisy = wb_daisy['Monday']
    daisy_row_count = sum(1 for row in monday_daisy.iter_rows() if row[0].value and row[0].value.strip())

    # Filtering to one arm removes exactly one data row (one fewer class-arm).
    assert daisy_row_count == all_row_count - 1


def test_export_by_day_pdf_respects_arm_filter(app):
    batch_id = _seed_print_filter_batch(app, 'G')
    c = _admin(app)

    r_all = c.get(f'/generator/results/{batch_id}/export_by_day_pdf')
    assert r_all.status_code == 200
    assert r_all.mimetype == 'application/pdf'

    r_daisy = c.get(f'/generator/results/{batch_id}/export_by_day_pdf?arm=ZzSSS2G|ZzDaisyG')
    assert r_daisy.status_code == 200
    assert r_daisy.mimetype == 'application/pdf'

    # A filtered-out selection with no match falls back to the batch's view page.
    r_none = c.get(f'/generator/results/{batch_id}/export_by_day_pdf?arm=NotAClass|NotAnArm')
    assert r_none.status_code == 302


def test_export_image_respects_arm_filter(app):
    from PIL import Image
    from io import BytesIO

    batch_id = _seed_print_filter_batch(app, 'H')
    c = _admin(app)

    r_all = c.get(f'/generator/results/{batch_id}/export_image_hd')
    assert r_all.status_code == 200
    assert r_all.mimetype == 'image/png'
    img_all = Image.open(BytesIO(r_all.data))
    all_height = img_all.height

    r_daisy = c.get(f'/generator/results/{batch_id}/export_image_hd?arm=ZzSSS2H|ZzDaisyH')
    assert r_daisy.status_code == 200
    assert r_daisy.mimetype == 'image/png'
    img_daisy = Image.open(BytesIO(r_daisy.data))
    # One fewer class-arm row means a shorter image (same width, less height).
    assert img_daisy.height < all_height
    assert img_daisy.width == img_all.width

    # A filtered-out selection with no match behaves like "no results".
    r_none = c.get(f'/generator/results/{batch_id}/export_image_hd?arm=NotAClass|NotAnArm')
    assert r_none.status_code == 302


def _seed_coscheduled_pair(app, tag):
    """A combined class with an ACTIVE co-schedule rule and matching
    GenTimetableResult rows on both sides of the same slot — the scenario
    where dropping one arm from a print/export should still surface its
    subject via the other arm's cell (e.g. "ZzLiterature/ZzAccounting")."""
    cc_id, lit_id, acct_id = _build_combined_class(app, tag)
    batch_id = f'zzbatch-pair-annotate-{tag}'
    with app.app_context():
        bid = Branch.get_default().id
        db.session.add(GenCoScheduleRule(
            branch_id=bid, name=f'ZzPairRule{tag}', source_subject_id=lit_id,
            source_class_name=f'ZzSSS2{tag}', source_arm_name=f'ZzDaisy{tag}',
            target_subject_id=acct_id, target_class_name=f'ZzSSS2{tag}', target_arm_name=f'ZzIris{tag}',
            is_active=True))
        db.session.add_all([
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name=f'ZzSSS2{tag}', arm_name=f'ZzDaisy{tag}', day_of_week=0,
                               period_number=1, subject_id=lit_id),
            GenTimetableResult(branch_id=bid, batch_id=batch_id, school_level='sss',
                               class_name=f'ZzSSS2{tag}', arm_name=f'ZzIris{tag}', day_of_week=0,
                               period_number=1, subject_id=acct_id),
        ])
        db.session.commit()
    return batch_id


def test_print_shows_coschedule_pair_only_when_other_arm_excluded(app):
    batch_id = _seed_coscheduled_pair(app, 'J')
    c = _admin(app)

    # Both arms included: Daisy's own cell should NOT carry the annotation —
    # Iris already has its own row showing Accounting.
    r_all = c.get(f'/generator/results/{batch_id}/print')
    assert r_all.status_code == 200
    body_all = r_all.get_data(as_text=True)
    assert 'ZzLiteratureJ' in body_all and 'ZzAccountingJ' in body_all
    assert 'ZzLiteratureJ/ZzAccountingJ' not in body_all

    # Iris excluded: Daisy's cell should now show both subjects.
    r_filtered = c.get(f'/generator/results/{batch_id}/print?arm=ZzSSS2J|ZzDaisyJ')
    assert r_filtered.status_code == 200
    body_filtered = r_filtered.get_data(as_text=True)
    assert 'ZzIrisJ' not in body_filtered   # the excluded arm's own row is gone
    assert 'ZzLiteratureJ/ZzAccountingJ' in body_filtered


def test_print_single_timetable_shows_coschedule_pair(app):
    """A single-arm print always excludes every other arm by definition, so
    the counterpart should always be annotated when a rule pairs it."""
    batch_id = _seed_coscheduled_pair(app, 'K')
    c = _admin(app)

    r = c.get(f'/generator/results/{batch_id}/print/ZzSSS2K/ZzDaisyK')
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert 'ZzLiteratureK/ZzAccountingK' in body


def test_export_by_class_shows_coschedule_pair(app):
    import openpyxl
    from io import BytesIO

    batch_id = _seed_coscheduled_pair(app, 'L')
    c = _admin(app)

    r = c.get(f'/generator/results/{batch_id}/export?arm=ZzSSS2L|ZzDaisyL')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]
    all_text = ' '.join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert '/' in all_text
    assert any(v.startswith('ZzLite') and '/ZzAcco' in v
               for row in ws.iter_rows() for v in [str(c.value) for c in row] if v and '/' in v)


def test_export_by_day_shows_coschedule_pair(app):
    import openpyxl
    from io import BytesIO

    batch_id = _seed_coscheduled_pair(app, 'M')
    c = _admin(app)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?arm=ZzSSS2M|ZzDaisyM')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    monday = wb['Monday']
    all_text = ' '.join(str(cell.value) for row in monday.iter_rows() for cell in row if cell.value)
    assert 'ZzIrisM' not in all_text   # excluded arm's own row is gone
    # export_by_day truncates to 5 chars (no abbrev_map match for these test names).
    assert any(v.startswith('ZzLit') and '/ZzAcc' in v
               for row in monday.iter_rows() for v in [str(c.value) for c in row] if v and '/' in v)


def test_export_by_day_pdf_shows_coschedule_pair(app):
    """PDF text isn't easily extractable in this test env, so this checks
    the route succeeds with the filter+annotation active together (the same
    code path exercised visually via Playwright) rather than the rendered
    text itself."""
    batch_id = _seed_coscheduled_pair(app, 'N')
    c = _admin(app)

    r = c.get(f'/generator/results/{batch_id}/export_by_day_pdf?arm=ZzSSS2N|ZzDaisyN')
    assert r.status_code == 200
    assert r.mimetype == 'application/pdf'
    assert len(r.data) > 500


def test_master_timetable_shows_coschedule_pair_when_other_arm_excluded(app):
    batch_id = _seed_coscheduled_pair(app, 'O')
    c = _admin(app)

    # Both arms included (default, no ?arm= at all): each keeps its own
    # subject, no annotation needed since the other arm already has a row.
    r_all = c.get(f'/generator/master-timetable?batch_id={batch_id}&day=0&period=1')
    assert r_all.status_code == 200
    body_all = r_all.get_data(as_text=True)
    assert 'ZzIrisO' in body_all   # Iris still has its own row
    assert 'ZzLiteratureO / ZzAccountingO' not in body_all

    # Daisy only: Daisy's row should now show both subjects, Iris's own
    # results row is gone (it still legitimately appears in the checkbox
    # panel, just unchecked, so check the results table specifically).
    r_daisy = c.get(f'/generator/master-timetable?batch_id={batch_id}&day=0&period=1&arm=ZzSSS2O|ZzDaisyO')
    assert r_daisy.status_code == 200
    body_daisy = r_daisy.get_data(as_text=True)
    assert '<strong>ZzSSS2O ZzDaisyO</strong>' in body_daisy
    assert '<strong>ZzSSS2O ZzIrisO</strong>' not in body_daisy
    assert 'ZzLiteratureO / ZzAccountingO' in body_daisy

    # The checkbox panel lists both arms regardless of which slot is viewed.
    assert 'name="arm" value="ZzSSS2O|ZzDaisyO"' in body_daisy
    assert 'name="arm" value="ZzSSS2O|ZzIrisO"' in body_daisy
