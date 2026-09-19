"""Subject-Wise Performance & Grade Breakdown report, independently for
WAEC, JAMB, Mock WAEC and Mock JAMB — per-subject, per-branch grade/score-band
percentages plus an overall summary by branch."""
import datetime
import uuid

from config import Config
from models import db, Student, Branch, WAECResult, JAMBResult, AcademicSession
from models.mock_waec import MockWAECExam, MockWAECResult
from models.mock_jamb import MockJAMBExam, MockJAMBResult
from tests.conftest import login_token


def _admin(app):
    c = app.test_client()
    c.post('/login', data={'password': Config.ADMIN_PASSWORD, '_csrf_token': login_token(c)})
    return c


def _student(branch_id, tag):
    s = Student(student_id='GB' + uuid.uuid4().hex[:7].upper(), first_name='Br',
                surname=tag, gender='Female', branch_id=branch_id)
    db.session.add(s)
    return s


def _branch(name):
    b = Branch.query.filter_by(name=name).first()
    if not b:
        b = Branch(name=name, code=uuid.uuid4().hex[:8].upper())
        db.session.add(b); db.session.flush()
    return b


def _active_session():
    sess = AcademicSession.query.filter_by(is_active=True).first()
    if not sess:
        sess = AcademicSession(name='2081/2082', is_active=True)
        db.session.add(sess); db.session.flush()
    return sess


def test_waec_breakdown_groups_by_subject_and_branch_with_not_offered(app):
    yr = 2081
    with app.app_context():
        b1 = _branch('GB Alpha')
        b2 = _branch('GB Beta')
        db.session.flush()
        s1 = _student(b1.id, 'One'); s2 = _student(b1.id, 'Two'); s3 = _student(b2.id, 'Three')
        db.session.flush()
        db.session.add_all([
            WAECResult(student_id=s1.id, exam_year=yr, subject='Mathematics', grade='A1'),
            WAECResult(student_id=s2.id, exam_year=yr, subject='Mathematics', grade='F9'),
            WAECResult(student_id=s3.id, exam_year=yr, subject='English Language', grade='C4'),
        ])
        db.session.commit()

    c = _admin(app)
    html = c.get(f'/results/subject-branch-breakdown?exam=waec&year={yr}').get_data(as_text=True)
    assert 'GB Alpha' in html and 'GB Beta' in html
    assert 'Mathematics' in html and 'English Language' in html
    # Beta never sat Mathematics -> "Not Offered" for that subject/branch pairing
    assert 'Not Offered' in html
    assert 'Overall Credit Pass (A1' in html


def test_jamb_breakdown_uses_score_bands(app):
    yr = 2082
    with app.app_context():
        b1 = _branch('GB Gamma')
        db.session.flush()
        s1 = _student(b1.id, 'Four')
        db.session.flush()
        db.session.add(JAMBResult(student_id=s1.id, exam_year=yr, total_score=180,
                                  subject1='English', subject1_score=92,
                                  subject2='Mathematics', subject2_score=18))
        db.session.commit()

    c = _admin(app)
    html = c.get(f'/results/subject-branch-breakdown?exam=jamb&year={yr}').get_data(as_text=True)
    assert 'Score Band' in html
    assert '90-100' in html and '0-19' in html
    assert 'GB Gamma' in html
    assert 'Overall ≥50 Rate' in html


def test_mock_waec_breakdown_selects_current_session_sitting(app):
    with app.app_context():
        sess = _active_session()
        b1 = _branch('GB Delta')
        db.session.flush()
        ex = MockWAECExam(name='GB Mock', exam_number=7, session_id=sess.id,
                          exam_date=datetime.date.today(), branch_id=b1.id)
        db.session.add(ex); db.session.flush()
        s1 = _student(b1.id, 'Five')
        db.session.flush()
        db.session.add(MockWAECResult(student_id=s1.id, mock_exam_id=ex.id,
                                      subject='Chemistry', score=80, grade='B2'))
        db.session.commit()

    c = _admin(app)
    html = c.get('/results/subject-branch-breakdown?exam=mock_waec').get_data(as_text=True)
    assert 'GB Delta' in html
    assert 'Chemistry' in html
    assert f'Mock 7' in html


def test_mock_jamb_breakdown_pivots_subject_slots(app):
    with app.app_context():
        sess_id = _active_session().id
        b1 = _branch('GB Epsilon')
        db.session.flush()
        ex = MockJAMBExam(name='GB JMock', exam_number=9, session_id=sess_id,
                          exam_date=datetime.date.today(), branch_id=b1.id)
        db.session.add(ex); db.session.flush()
        s1 = _student(b1.id, 'Six')
        db.session.flush()
        db.session.add(MockJAMBResult(student_id=s1.id, mock_exam_id=ex.id, total_score=140,
                                      subject1='Physics', subject1_score=55,
                                      subject2='Biology', subject2_score=40))
        db.session.commit()

    c = _admin(app)
    html = c.get(f'/results/subject-branch-breakdown?exam=mock_jamb&mock={sess_id}_9').get_data(as_text=True)
    assert 'GB Epsilon' in html
    assert 'Physics' in html and 'Biology' in html


def test_no_data_shows_empty_state(app):
    c = _admin(app)
    html = c.get('/results/subject-branch-breakdown?exam=waec&year=1904').get_data(as_text=True)
    assert 'No WAEC results recorded' in html


def test_exam_type_switcher_links_present(app):
    c = _admin(app)
    html = c.get('/results/subject-branch-breakdown').get_data(as_text=True)
    assert 'exam=waec' in html and 'exam=jamb' in html
    assert 'exam=mock_waec' in html and 'exam=mock_jamb' in html


def test_download_buttons_present_when_data_exists(app):
    yr = 2084
    with app.app_context():
        b1 = _branch('GB Zeta')
        db.session.flush()
        s1 = _student(b1.id, 'Seven')
        db.session.flush()
        db.session.add(WAECResult(student_id=s1.id, exam_year=yr, subject='Mathematics', grade='A1'))
        db.session.commit()

    c = _admin(app)
    html = c.get(f'/results/subject-branch-breakdown?exam=waec&year={yr}').get_data(as_text=True)
    assert '/subject-branch-breakdown/export.pdf' in html
    assert '/subject-branch-breakdown/export.docx' in html
    assert '/subject-branch-breakdown/export.xlsx' in html
    assert '/subject-branch-breakdown/export.png' in html


def test_export_without_data_redirects_with_flash(app):
    c = _admin(app)
    r = c.get('/results/subject-branch-breakdown/export.pdf?exam=waec&year=1904')
    assert r.status_code == 302
    assert '/subject-branch-breakdown' in r.headers['Location']


def test_export_rejects_unknown_format(app):
    c = _admin(app)
    r = c.get('/results/subject-branch-breakdown/export.bogus?exam=waec')
    assert r.status_code == 404


def test_waec_export_pdf_docx_xlsx_png(app):
    yr = 2085
    with app.app_context():
        b1 = _branch('GB Eta')
        b2 = _branch('GB Theta')
        db.session.flush()
        s1 = _student(b1.id, 'Eight'); s2 = _student(b2.id, 'Nine')
        db.session.flush()
        db.session.add_all([
            WAECResult(student_id=s1.id, exam_year=yr, subject='Mathematics', grade='A1'),
            WAECResult(student_id=s2.id, exam_year=yr, subject='English Language', grade='C6'),
        ])
        db.session.commit()

    c = _admin(app)
    r = c.get(f'/results/subject-branch-breakdown/export.pdf?exam=waec&year={yr}')
    assert r.status_code == 200 and r.mimetype == 'application/pdf'
    assert r.get_data()[:4] == b'%PDF'

    r = c.get(f'/results/subject-branch-breakdown/export.docx?exam=waec&year={yr}')
    assert r.status_code == 200
    assert r.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    from docx import Document
    import io
    doc = Document(io.BytesIO(r.get_data()))
    # the navy masthead is itself a (1-column) table, so filter down to the
    # two real data tables: main breakdown + overall summary
    data_tables = [t for t in doc.tables if len(t.columns) > 1]
    assert len(data_tables) == 2
    header = [c_.text for c_ in data_tables[0].rows[0].cells]
    assert header[:3] == ['SUBJECT', 'BRANCH', 'CANDIDATES (N)']
    band_header = [c_.text for c_ in data_tables[0].rows[1].cells]
    assert 'A1 (%)' in band_header

    r = c.get(f'/results/subject-branch-breakdown/export.xlsx?exam=waec&year={yr}')
    assert r.status_code == 200
    assert r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(r.get_data()))
    assert wb.sheetnames == ['Grade Breakdown', 'Summary']

    r = c.get(f'/results/subject-branch-breakdown/export.png?exam=waec&year={yr}')
    assert r.status_code == 200
    assert r.mimetype in ('image/png', 'application/zip')


def test_jamb_export_uses_score_bands(app):
    yr = 2086
    with app.app_context():
        b1 = _branch('GB Iota')
        db.session.flush()
        s1 = _student(b1.id, 'Ten')
        db.session.flush()
        db.session.add(JAMBResult(student_id=s1.id, exam_year=yr, total_score=180,
                                  subject1='English', subject1_score=92))
        db.session.commit()

    c = _admin(app)
    r = c.get(f'/results/subject-branch-breakdown/export.xlsx?exam=jamb&year={yr}')
    assert r.status_code == 200
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(r.get_data()))
    ws = wb['Grade Breakdown']
    all_values = [c.value for row in ws.iter_rows() for c in row]
    assert any(v and '90-100' in str(v) for v in all_values)
