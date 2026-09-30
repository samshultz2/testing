"""Site-wide "time format" setting (12h/24h) for how period and school-day
clock times are displayed: the timetable generator's settings preview, its
PDF/XLSX/image exports, and the published timetable module (SPA + print
PDF). A single SchoolSettings key ('time_format'), same storage pattern as
the existing timezone setting -- no schema change, so no migration."""
from io import BytesIO
from datetime import time as dtime

from config import Config
from models import db, Branch, GenTimetableRule, GenTimetableResult, SchoolSettings, TimetableSlot
from tests.conftest import login_token


def _admin(app):
    c = app.test_client()
    c.post('/login', data={'password': Config.ADMIN_PASSWORD, '_csrf_token': login_token(c)})
    with c.session_transaction() as s:
        s['_csrf_token'] = 'a' * 64
    return c


def _post(c, url, **data):
    data.setdefault('_csrf_token', 'a' * 64)
    return c.post(url, data=data)


def _scoped_to_branch(c, branch_id):
    with c.session_transaction() as s:
        s['view_branch_id'] = branch_id
    return c


def _set_time_format(app, value):
    with app.app_context():
        SchoolSettings.set('time_format', value, 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()


# --- utils.timeutil.format_clock / get_time_format --------------------------

def test_format_clock_24h_pads_hour_and_minute():
    from utils.timeutil import format_clock
    assert format_clock(8, 5, '24h') == '08:05'
    assert format_clock(0, 0, '24h') == '00:00'
    assert format_clock(13, 40, '24h') == '13:40'
    assert format_clock(23, 59, '24h') == '23:59'


def test_format_clock_12h_no_leading_zero_and_meridiem():
    from utils.timeutil import format_clock
    assert format_clock(0, 0, '12h') == '12:00 AM'      # midnight
    assert format_clock(8, 5, '12h') == '8:05 AM'
    assert format_clock(12, 0, '12h') == '12:00 PM'      # noon
    assert format_clock(13, 40, '12h') == '1:40 PM'
    assert format_clock(23, 59, '12h') == '11:59 PM'


def test_format_clock_wraps_hours_past_24():
    # exports.py accumulates minutes across periods/breaks and can legitimately
    # roll past midnight for a very long or late-starting school day.
    from utils.timeutil import format_clock
    assert format_clock(25, 5, '24h') == '01:05'
    assert format_clock(25, 5, '12h') == '1:05 AM'


def test_get_time_format_defaults_to_24h_and_respects_setting(app):
    from utils.timeutil import get_time_format, clear_time_format_cache
    with app.app_context():
        SchoolSettings.query.filter_by(key='time_format').delete()
        db.session.commit()
        clear_time_format_cache()
        assert get_time_format() == '24h'

    _set_time_format(app, '12h')
    with app.app_context():
        assert get_time_format() == '12h'

    _set_time_format(app, '24h')
    with app.app_context():
        assert get_time_format() == '24h'


def test_get_time_format_ignores_garbage_value(app):
    with app.app_context():
        SchoolSettings.set('time_format', 'not-a-format', 'string', 'test')
        from utils.timeutil import clear_time_format_cache, get_time_format
        clear_time_format_cache()
        assert get_time_format() == '24h'


# --- utils.generator_times.day_end_time --------------------------------------

def test_day_end_time_respects_format(app):
    from utils.generator_times import day_end_time
    rules = {'day_start': '8:20', 'period_minutes': '40', 'break_minutes': '30'}
    _set_time_format(app, '24h')
    with app.app_context():
        assert day_end_time(rules, periods_per_day=8, break_after=4) == '14:10'
    _set_time_format(app, '12h')
    with app.app_context():
        assert day_end_time(rules, periods_per_day=8, break_after=4) == '2:10 PM'


# --- Settings page: persists and reflects the setting ------------------------

def test_school_settings_saves_and_reflects_time_format(app):
    c = _admin(app)
    r = _post(c, '/settings/school', school_name='ZzTFSchool', time_format='12h')
    assert r.status_code in (200, 302)
    with app.app_context():
        assert SchoolSettings.get('time_format') == '12h'

    body = c.get('/settings/school').get_data(as_text=True)
    assert '"current_time_format": "12h"' in body or '\\"current_time_format\\":\\"12h\\"' in body \
        or '12h' in body   # SPA JSON payload is embedded either escaped or not

    _post(c, '/settings/school', school_name='ZzTFSchool', time_format='24h')
    with app.app_context():
        assert SchoolSettings.get('time_format') == '24h'


def test_school_settings_ignores_invalid_time_format(app):
    _set_time_format(app, '12h')
    c = _admin(app)
    _post(c, '/settings/school', school_name='ZzTFSchool2', time_format='bogus')
    with app.app_context():
        # unchanged -- an invalid value is silently ignored, not stored.
        assert SchoolSettings.get('time_format') == '12h'


# --- Generator exports: XLSX period-time headers respect the format ----------

def _seed_generator_batch(app, tag):
    with app.app_context():
        b = Branch(name=f'ZzTFBranch{tag}', code=None)
        db.session.add(b); db.session.flush()
        bid = b.id
        batch_id = f'zztf-{tag}'
        db.session.add_all([
            GenTimetableRule(branch_id=bid, rule_type='periods_per_day', value='8',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_after_period', value='4',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='day_start', value='8:20',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='period_minutes', value='40',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_minutes', value='30',
                             school_level='sss', is_active=True),
        ])
        for d in range(5):
            for p in (1, 2):
                db.session.add(GenTimetableResult(
                    branch_id=bid, batch_id=batch_id, school_level='sss',
                    class_name=f'ZzTF{tag}', arm_name='Main',
                    day_of_week=d, period_number=p))
        db.session.commit()
        return batch_id, bid


def test_export_by_class_xlsx_period_headers_use_configured_format(app):
    import openpyxl
    batch_id, bid = _seed_generator_batch(app, 'XL')
    c = _scoped_to_branch(_admin(app), bid)

    _set_time_format(app, '24h')
    r24 = c.get(f'/generator/results/{batch_id}/export')
    assert r24.status_code == 200
    wb24 = openpyxl.load_workbook(BytesIO(r24.data))
    text24 = ' '.join(str(cell.value) for row in wb24[wb24.sheetnames[0]].iter_rows()
                      for cell in row if cell.value)
    assert '8:20' in text24 or '08:20' in text24
    assert 'AM' not in text24 and 'PM' not in text24

    _set_time_format(app, '12h')
    r12 = c.get(f'/generator/results/{batch_id}/export')
    assert r12.status_code == 200
    wb12 = openpyxl.load_workbook(BytesIO(r12.data))
    text12 = ' '.join(str(cell.value) for row in wb12[wb12.sheetnames[0]].iter_rows()
                      for cell in row if cell.value)
    assert 'AM' in text12 or 'PM' in text12
    assert '8:20 AM' in text12


def test_export_by_day_xlsx_period_headers_use_configured_format(app):
    import openpyxl
    batch_id, bid = _seed_generator_batch(app, 'BD')
    c = _scoped_to_branch(_admin(app), bid)

    _set_time_format(app, '12h')
    r = c.get(f'/generator/results/{batch_id}/export_by_day')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    text = ' '.join(str(cell.value) for name in wb.sheetnames
                    for row in wb[name].iter_rows() for cell in row if cell.value)
    assert 'AM' in text or 'PM' in text


# --- Published timetable module: SPA slot dict + print PDF -------------------

def test_slot_dict_uses_configured_time_format(app):
    from routes.timetable import _slot_dict
    slot = TimetableSlot(slot_number=1, name='Period 1', start_time=dtime(8, 20),
                         end_time=dtime(9, 0), is_break=False, order=1, is_active=True)

    _set_time_format(app, '24h')
    with app.app_context():
        d = _slot_dict(slot)
        assert d['start'] == '08:20' and d['end'] == '09:00'

    _set_time_format(app, '12h')
    with app.app_context():
        d = _slot_dict(slot)
        assert d['start'] == '8:20 AM' and d['end'] == '9:00 AM'
