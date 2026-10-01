"""Two display fixes for the timetable generator's exports:

1. The by-day PNG image export cramped the period header's clock range onto
   one line ("8:00 AM-8:40 AM"), which overlapped the next column's text
   under 12h format. Now it's 3 stacked lines (period / start / end).
2. The A3 "packed" (2-days-per-page) XLSX export repeated a full
   "Class | P1...P9 | BREAK | ..." header row for the second day of each
   pair, even though periods are identical every day. Now only the first
   day of a pair gets that row; the second's data starts right under its
   own day-name bar."""
from io import BytesIO

import pytest
from config import Config
from models import (
    db, Branch, GenSubject, GenTeacher, GenTeacherAssignment, GenClassConfig,
    GenClassSubjectConfig, GenTimetableRule, GenTimetableResult, SchoolSettings,
)
from tests.conftest import login_token


@pytest.fixture(autouse=True)
def _restore_time_format(app):
    """time_format is a site-wide SchoolSettings row (not branch-scoped), so
    unlike everything else in this file it isn't isolated by a dedicated
    branch -- every test here sets it, so every test here must put it back,
    or a later file's test that assumes the default ('24h') can fail
    depending on alphabetical test-file ordering."""
    yield
    with app.app_context():
        SchoolSettings.set('time_format', '24h', 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()


def _admin(app):
    c = app.test_client()
    c.post('/login', data={'password': Config.ADMIN_PASSWORD, '_csrf_token': login_token(c)})
    with c.session_transaction() as s:
        s['_csrf_token'] = 'a' * 64
    return c


def _scoped_to_branch(c, branch_id):
    with c.session_transaction() as s:
        s['view_branch_id'] = branch_id
    return c


def _seed(app, tag, periods_per_day=9, break_after=5):
    """One class, one arm, a handful of periods across all 5 days -- seeded
    into a DEDICATED branch so periods_per_day/break_after/day clock settings
    (all per-branch) can't be polluted by other test files sharing the
    default branch. Returns (batch_id, branch_id)."""
    with app.app_context():
        b = Branch(name=f'ZzExportTimeBranch{tag}', code=None)
        db.session.add(b); db.session.flush()
        bid = b.id

        db.session.add_all([
            GenTimetableRule(branch_id=bid, rule_type='periods_per_day', value=str(periods_per_day),
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_after_period', value=str(break_after),
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='day_start', value='8:00',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='period_minutes', value='40',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_minutes', value='30',
                             school_level='sss', is_active=True),
        ])
        SchoolSettings.set('time_format', '12h', 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()

        cc = GenClassConfig(branch_id=bid, class_name=f'ZzET{tag}', school_level='sss',
                            num_arms=1, arm_names=f'ZzArm{tag}', has_streams=False)
        db.session.add(cc); db.session.flush()

        subj = GenSubject(branch_id=bid, name=f'ZzSubj{tag}', school_level='sss')
        db.session.add(subj); db.session.flush()

        t = GenTeacher(branch_id=bid, name=f'Zz Teacher {tag}', school_level='sss',
                       max_periods_per_day=periods_per_day, max_periods_per_week=periods_per_day * 5)
        db.session.add(t); db.session.flush()
        db.session.add(GenTeacherAssignment(branch_id=bid, teacher_id=t.id, subject_id=subj.id,
                                            class_config_id=cc.id, arm_name=f'ZzArm{tag}'))
        db.session.add(GenClassSubjectConfig(
            class_config_id=cc.id, subject_id=subj.id, is_enabled=True, periods_per_week=5))

        batch_id = f'zzettime-{tag}'
        for d in range(5):
            db.session.add(GenTimetableResult(
                branch_id=bid, batch_id=batch_id, school_level='sss',
                class_name=f'ZzET{tag}', arm_name=f'ZzArm{tag}',
                day_of_week=d, period_number=1, subject_id=subj.id, teacher_id=t.id))
        db.session.commit()
        return batch_id, bid


# --- Image export: stacked period header -------------------------------------

def test_image_export_header_height_reflects_three_stacked_lines(app):
    """The image's total height is a direct function of header_height (one
    term per day block) -- pins that constant at 70*scale (3 lines) rather
    than the old 55*scale (2 lines), so a regression back to the cramped
    2-line format would change this and fail loudly."""
    from PIL import Image
    batch_id, bid = _seed(app, 'IMG', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_image_hd')
    assert r.status_code == 200
    assert r.mimetype == 'image/png'
    img = Image.open(BytesIO(r.data))

    scale = 4   # 'hd' quality
    cell_height = 35 * scale
    header_height = 70 * scale
    day_header_height = 55 * scale
    margin = 50 * scale
    day_spacing = 25 * scale
    title_height = 120 * scale
    num_class_arms = 1
    num_days = 5

    expected_height = title_height + (num_days * (day_header_height + header_height
                      + (num_class_arms * cell_height) + day_spacing)) + (margin * 2)
    assert img.height == expected_height


def test_image_export_break_column_widened_to_fit_break_time(app):
    """break_col_width went from 35*scale to 48*scale so the stacked break
    start/end time ("11:20 AM" etc.) actually fits instead of being clipped."""
    from PIL import Image
    batch_id, bid = _seed(app, 'IMGW', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_image_hd')
    assert r.status_code == 200
    img = Image.open(BytesIO(r.data))

    scale = 4
    cell_width = 85 * scale
    class_col_width = 55 * scale
    break_col_width = 48 * scale
    margin = 50 * scale
    periods_per_day = 9
    table_width = class_col_width + (periods_per_day * cell_width) + break_col_width
    expected_width = table_width + (margin * 2)
    assert img.width == expected_width


def test_image_export_renders_without_error_in_24h_format_too(app):
    """24h format never had the overlap bug (no AM/PM suffix), but the
    3-line layout change must not break it either."""
    from PIL import Image
    batch_id, bid = _seed(app, 'IMG24', periods_per_day=9, break_after=5)
    with app.app_context():
        SchoolSettings.set('time_format', '24h', 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_image_hd')
    assert r.status_code == 200
    img = Image.open(BytesIO(r.data))
    assert img.width > 0 and img.height > 0


# --- A3 packed XLSX export: no repeated header for the 2nd day ---------------

def test_packed_export_second_day_has_no_repeated_header_row(app):
    import openpyxl
    batch_id, bid = _seed(app, 'PK', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]   # Monday & Tuesday packed onto one sheet

    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=12)]
    # Row 1: school name. Row 2: MONDAY bar. Row 3: the one-and-only
    # "Class | P1...P9 | BREAK | ..." header, WITH times (first day overall).
    assert rows[1][0] == 'MONDAY'
    assert rows[2][0] == 'Class'
    assert 'AM' in rows[2][1] or 'PM' in rows[2][1]   # has clock times

    # Monday's single data row (one class-arm) comes right after.
    assert rows[3][0] == 'ZzETPKZ'

    # A blank spacer row, then TUESDAY's bar...
    tuesday_row_idx = next(i for i, r in enumerate(rows) if r[0] == 'TUESDAY')
    # ...immediately followed by Tuesday's DATA row -- not another "Class" header.
    next_row = rows[tuesday_row_idx + 1]
    assert next_row[0] == 'ZzETPKZ'
    assert next_row[0] != 'Class'
    assert not any(isinstance(v, str) and v.startswith('P1') for v in next_row)


def test_packed_export_wed_thu_pair_also_has_single_header(app):
    """Confirms the fix isn't somehow tied to day_of_week==0 (Monday)
    specifically -- the SAME rule (header only on the first day of each
    packed pair) must hold for every pair, including one that doesn't
    include Monday at all."""
    import openpyxl
    batch_id, bid = _seed(app, 'PK2', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    assert 'Wed & Thu' in wb.sheetnames
    ws = wb['Wed & Thu']

    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=10)]
    assert rows[0][0] == 'WEDNESDAY'
    assert rows[1][0] == 'Class'        # Wednesday still gets its header (first of the pair)
    assert rows[1][1] == 'P1'           # no times -- Wednesday isn't the global day 0
    assert rows[2][0] == 'ZzETPK2Z'

    thursday_row_idx = next(i for i, r in enumerate(rows) if r[0] == 'THURSDAY')
    next_row = rows[thursday_row_idx + 1]
    assert next_row[0] == 'ZzETPK2Z'    # Thursday's data, no repeated header row


def test_single_day_per_page_export_unaffected(app):
    """The non-packed (1 day/page) export must still show a header on
    EVERY day -- the "skip the 2nd day's header" logic is specific to
    packed pairs, not a change to single-day pages."""
    import openpyxl
    batch_id, bid = _seed(app, 'SNG', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day')  # default layout=single
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    assert len(wb.sheetnames) == 5
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=4)]
        header_row = next(r for r in rows if r[0] == 'Class')
        assert header_row is not None
