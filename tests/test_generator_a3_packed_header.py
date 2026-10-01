"""The A3 "packed" (2-days-per-page) XLSX export:

1. Repeated a full "Class | P1...P9 | BREAK | ..." header row for the
   second day of each pair, even though periods are identical every day.
   Now only the first day of a pair gets that row; the second's data
   starts right under its own day-name bar. Single-day-per-page exports
   are unaffected.
2. The period header's clock range was on one line ("8:00 AM-8:40 AM").
   Now it's 3 stacked lines (period / start / end).
3. Data-row height used to differ between a day sharing its page with the
   school name/address (Monday) and every other day, and the "BREAK"
   column index was hardcoded to 7 (wrong for any break_after != 5). Both
   are now computed once and applied identically on every day, every page.

No colors or styling changed -- structural/sizing only."""
from io import BytesIO

import pytest
from config import Config
from models import (
    db, Branch, GenSubject, GenTeacher, GenTeacherAssignment, GenClassConfig,
    GenClassSubjectConfig, GenTimetableRule, GenTimetableResult, SchoolSettings,
)
from tests.conftest import login_token


@pytest.fixture(autouse=True)
def _restore_time_format(app):
    """time_format is a site-wide SchoolSettings row (not branch-scoped), so
    every test here that sets it must put it back for other test files."""
    yield
    with app.app_context():
        SchoolSettings.set('time_format', '24h', 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()


def _admin(app):
    c = app.test_client()
    c.post('/login', data={'password': Config.ADMIN_PASSWORD, '_csrf_token': login_token(c)})
    with c.session_transaction() as s:
        s['_csrf_token'] = 'a' * 64
    return c


def _scoped_to_branch(c, branch_id):
    with c.session_transaction() as s:
        s['view_branch_id'] = branch_id
    return c


def _seed(app, tag, periods_per_day=9, break_after=5):
    with app.app_context():
        b = Branch(name=f'ZzA3HeaderBranch{tag}', code=None)
        db.session.add(b); db.session.flush()
        bid = b.id

        db.session.add_all([
            GenTimetableRule(branch_id=bid, rule_type='periods_per_day', value=str(periods_per_day),
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_after_period', value=str(break_after),
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='day_start', value='8:00',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='period_minutes', value='40',
                             school_level='sss', is_active=True),
            GenTimetableRule(branch_id=bid, rule_type='break_minutes', value='30',
                             school_level='sss', is_active=True),
        ])
        SchoolSettings.set('time_format', '12h', 'string', 'test')
        from utils.timeutil import clear_time_format_cache
        clear_time_format_cache()

        cc = GenClassConfig(branch_id=bid, class_name=f'ZzA3{tag}', school_level='sss',
                            num_arms=1, arm_names=f'ZzArm{tag}', has_streams=False)
        db.session.add(cc); db.session.flush()

        subj = GenSubject(branch_id=bid, name=f'ZzSubj{tag}', school_level='sss')
        db.session.add(subj); db.session.flush()

        t = GenTeacher(branch_id=bid, name=f'Zz Teacher {tag}', school_level='sss',
                       max_periods_per_day=periods_per_day, max_periods_per_week=periods_per_day * 5)
        db.session.add(t); db.session.flush()
        db.session.add(GenTeacherAssignment(branch_id=bid, teacher_id=t.id, subject_id=subj.id,
                                            class_config_id=cc.id, arm_name=f'ZzArm{tag}'))
        db.session.add(GenClassSubjectConfig(
            class_config_id=cc.id, subject_id=subj.id, is_enabled=True, periods_per_week=5))

        batch_id = f'zza3hdr-{tag}'
        for d in range(5):
            db.session.add(GenTimetableResult(
                branch_id=bid, batch_id=batch_id, school_level='sss',
                class_name=f'ZzA3{tag}', arm_name=f'ZzArm{tag}',
                day_of_week=d, period_number=1, subject_id=subj.id, teacher_id=t.id))
        db.session.commit()
        return batch_id, bid


def test_packed_export_second_day_has_no_repeated_header_row(app):
    import openpyxl
    batch_id, bid = _seed(app, 'PK', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]   # Monday & Tuesday packed onto one sheet

    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=12)]
    # Row 1: school name. Row 2: MONDAY bar. Row 3: the one-and-only
    # "Class | P1...P9 | BREAK | ..." header, WITH times (first day overall).
    assert rows[1][0] == 'MONDAY'
    assert rows[2][0] == 'Class'
    assert 'AM' in rows[2][1] or 'PM' in rows[2][1]   # has clock times

    # Monday's single data row (one class-arm) comes right after.
    assert rows[3][0] == 'ZzA3PKZ'

    # A blank spacer row, then TUESDAY's bar...
    tuesday_row_idx = next(i for i, r in enumerate(rows) if r[0] == 'TUESDAY')
    # ...immediately followed by Tuesday's DATA row -- not another "Class" header.
    next_row = rows[tuesday_row_idx + 1]
    assert next_row[0] == 'ZzA3PKZ'
    assert next_row[0] != 'Class'
    assert not any(isinstance(v, str) and v.startswith('P1') for v in next_row)


def test_packed_export_wed_thu_pair_also_has_single_header(app):
    """Confirms the fix isn't somehow tied to day_of_week==0 (Monday)
    specifically -- the SAME rule (header only on the first day of each
    packed pair) must hold for every pair, including one that doesn't
    include Monday at all."""
    import openpyxl
    batch_id, bid = _seed(app, 'PK2', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    assert 'Wed & Thu' in wb.sheetnames
    ws = wb['Wed & Thu']

    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=10)]
    assert rows[0][0] == 'WEDNESDAY'
    assert rows[1][0] == 'Class'        # Wednesday still gets its header (first of the pair)
    assert rows[1][1] == 'P1'           # no times -- Wednesday isn't the global day 0
    assert rows[2][0] == 'ZzA3PK2Z'

    thursday_row_idx = next(i for i, r in enumerate(rows) if r[0] == 'THURSDAY')
    next_row = rows[thursday_row_idx + 1]
    assert next_row[0] == 'ZzA3PK2Z'    # Thursday's data, no repeated header row


def test_single_day_per_page_export_unaffected(app):
    """The non-packed (1 day/page) export must still show a header on
    EVERY day -- the "skip the 2nd day's header" logic is specific to
    packed pairs, not a change to single-day pages."""
    import openpyxl
    batch_id, bid = _seed(app, 'SNG', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day')  # default layout=single
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    assert len(wb.sheetnames) == 5
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=4)]
        header_row = next(r for r in rows if r[0] == 'Class')
        assert header_row is not None


def test_no_fill_colors_introduced(app):
    """This fix is structural only -- confirms no cell fill/color was added
    anywhere in the packed sheet (the real export has always been pure
    black-on-white for B&W printing)."""
    import openpyxl
    batch_id, bid = _seed(app, 'NOCOL', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            fg = cell.fill.fgColor
            # openpyxl's default "no fill" is patternType None / indexed 64 ('00000000' or theme-less).
            assert cell.fill.patternType is None, f'unexpected fill at {cell.coordinate}: {cell.fill}'


def test_period_header_time_is_stacked_three_lines(app):
    import openpyxl
    batch_id, bid = _seed(app, 'STACK', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]
    header_row = [c.value for c in list(ws.iter_rows(min_row=3, max_row=3))[0]]
    p1_cell = header_row[1]
    # "P1\n8:00 AM\n8:40 AM" -- 3 lines, not "P1\n8:00 AM-8:40 AM" (2 lines).
    assert p1_cell.count('\n') == 2
    lines = p1_cell.split('\n')
    assert lines[0] == 'P1'
    assert '-' not in lines[1] and '-' not in lines[2]
    assert 'AM' in lines[1] and 'AM' in lines[2]


def test_data_row_height_consistent_across_every_day_and_sheet(app):
    """Monday (which shares its page with the school name/address) must end
    up with the exact same data-row height as every other day, on every
    sheet of the workbook -- previously Monday's rows were visibly
    shorter."""
    import openpyxl
    batch_id, bid = _seed(app, 'ROWH', periods_per_day=9, break_after=5)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))

    heights = set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Data rows are the ones holding this test's single class-arm code.
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == 'ZzA3ROWHZ':
                    heights.add(ws.row_dimensions[cell.row].height)
    assert len(heights) == 1, f'data-row heights differ across days/sheets: {heights}'


def test_break_column_width_follows_actual_break_after(app):
    """The BREAK column's narrower width used to be hardcoded to column
    index 7 (only correct when break_after == 5) -- with a different
    break_after, that hardcoded index would narrow a period column
    instead and leave BREAK full-width. Uses break_after=4, which places
    BREAK at a different column (6, not 7)."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    batch_id, bid = _seed(app, 'BRKW', periods_per_day=9, break_after=4)
    c = _scoped_to_branch(_admin(app), bid)

    r = c.get(f'/generator/results/{batch_id}/export_by_day?paper=a3&layout=packed')
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb[wb.sheetnames[0]]

    header_row = [c.value for c in list(ws.iter_rows(min_row=3, max_row=3))[0]]
    # BREAK's value at d==0 is the break_time string (starts with "BREAK").
    break_col_idx = next(i for i, v in enumerate(header_row, start=1)
                        if isinstance(v, str) and v.startswith('BREAK'))
    assert break_col_idx == 6   # Class(1) + 4 before-break periods -> BREAK at 6

    narrow_width = ws.column_dimensions[get_column_letter(break_col_idx)].width
    wide_width = ws.column_dimensions[get_column_letter(break_col_idx - 1)].width
    assert narrow_width < wide_width
