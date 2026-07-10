"""
Communication routes — the school's communication center. Broadcast to parents
over SMS, WhatsApp (click-to-chat) or Email from one composer, with reusable
templates, audience targeting (class, arm, fee defaulters, selected students),
scheduling, a full send log with per-recipient delivery status, and CSV export.
Automated notifications from other modules funnel through the same campaign
engine (see utils.comms.build_campaign / create_draft_campaign).
"""
from datetime import datetime
from utils.helpers import get_active_term
import csv
import io

from flask import (Blueprint, request, redirect, url_for,
                   flash, jsonify, Response)
from sqlalchemy import func

from models import (
    db, Student, Term, SchoolClass, ClassArm, StudentEnrollment,
    ClassArmAssignment, MessageTemplate,
    Message, MessageRecipient, Announcement,
)
from utils.access_control import login_required, admin_required
from utils.branch_scope import scope_query, require_branch_access
from utils import comms
from utils.security import strip_tags

comms_bp = Blueprint('comms', __name__, url_prefix='/communication')

# Delivery channels the composer offers. SMS/WhatsApp reach a phone number; Email
# reaches a stored email address; In-app posts a bell notification to staff with a
# user account. (WhatsApp is a click-to-chat deep link; SMS/Email send through the
# configured gateway / SMTP; In-app is instant.)
CHANNELS = ['SMS', 'WhatsApp', 'Email', 'In-app']


# --- SPA helpers (no-reload React shell + JSON-aware action responses) ---
from utils.spa import section_responders
from utils.search import like_term
_wants_json, _render, _ok, _err = section_responders(
    'communication/app.html', 'comm_json', 'comms.dashboard')


def _is_admin():
    from utils.access_control import is_admin
    return is_admin()


def _nav_urls():
    return {'dashboard': url_for('comms.dashboard'), 'compose': url_for('comms.compose'),
            'inbox': url_for('comms.inbox'),
            'announcements': url_for('comms.announcements'), 'messages': url_for('comms.messages_list'),
            'reports': url_for('comms.reports'),
            'templates': url_for('comms.templates_list'), 'contacts': url_for('comms.contacts'),
            'settings': url_for('comms.settings')}


def _active_term():
    return get_active_term()


def _term_from(tid):
    """Term for an id (None-safe), falling back to the active term."""
    return (db.session.get(Term, tid) if tid else None) or _active_term()


def _dt(value):
    """Parse an HTML datetime-local value ('YYYY-MM-DDTHH:MM')."""
    for fmt in ('%Y-%m-%dT%H:%M', '%Y-%m-%d %H:%M'):
        try:
            return datetime.strptime(value, fmt)
        except (ValueError, TypeError):
            continue
    return None


# ============================================================================
# DASHBOARD
# ============================================================================

@comms_bp.route('/')
@login_required
def dashboard():
    from datetime import datetime, time, date
    cov = comms.coverage_stats()
    base_msgs = scope_query(Message.query, Message)
    total_campaigns = base_msgs.count()
    total_recipients = scope_query(db.session.query(func.coalesce(func.sum(Message.recipient_count), 0)), Message).scalar() or 0
    total_sent = scope_query(db.session.query(func.coalesce(func.sum(Message.sent_count), 0)), Message).scalar() or 0
    channel_rows = (scope_query(db.session.query(Message.channel, func.count(Message.id)), Message)
                    .group_by(Message.channel).all())
    channel_chart = [{'channel': ch or 'Other', 'count': n} for ch, n in channel_rows]
    template_count = MessageTemplate.query.filter_by(is_active=True).count()

    # Pipeline snapshot — what needs attention right now.
    scheduled_ct = base_msgs.filter(Message.status == 'Scheduled').count()
    draft_ct = base_msgs.filter(Message.status == 'Draft').count()

    # Recipient-level delivery stats (scoped by the owning campaign's branch).
    def _recip_q():
        return scope_query(
            db.session.query(MessageRecipient).join(Message, MessageRecipient.message_id == Message.id),
            Message)
    status_rows = (_recip_q().with_entities(MessageRecipient.status, func.count(MessageRecipient.id))
                   .group_by(MessageRecipient.status).all())
    by_status = {s: n for s, n in status_rows}
    sent_all = by_status.get('Sent', 0)
    failed_all = by_status.get('Failed', 0)
    attempted = sent_all + failed_all
    success_rate = round(sent_all / attempted * 100, 1) if attempted else None

    # Today's deliveries, split by channel (SMS / Email / other).
    start = datetime.combine(date.today(), time.min)
    today_rows = (_recip_q().with_entities(Message.channel, func.count(MessageRecipient.id))
                  .filter(MessageRecipient.status == 'Sent', MessageRecipient.sent_at >= start)
                  .group_by(Message.channel).all())
    today_by_channel = {(ch or 'Other'): n for ch, n in today_rows}

    def _chan(*names):
        return sum(today_by_channel.get(n, 0) for n in names)
    sent_today = sum(today_by_channel.values())

    recent = base_msgs.order_by(Message.created_at.desc()).limit(8).all()
    return _render({
        'page': 'dashboard', 'nav': _nav_urls(),
        'cov': cov, 'total_campaigns': total_campaigns,
        'total_recipients': int(total_recipients), 'total_sent': int(total_sent),
        'channel_chart': channel_chart, 'template_count': template_count,
        'stats': {
            'sent_today': sent_today,
            'sms_today': _chan('SMS', 'WhatsApp'),
            'email_today': _chan('Email'),
            'scheduled': scheduled_ct, 'drafts': draft_ct,
            'failed': failed_all, 'success_rate': success_rate,
        },
        'recent': [{'id': m.id, 'date': m.created_at.strftime('%d %b'),
                    'title': m.title or 'Message', 'audience_label': m.audience_label,
                    'channel': m.channel, 'status': m.status, 'sent_count': m.sent_count or 0,
                    'recipient_count': m.recipient_count or 0,
                    'url': url_for('comms.message_detail', message_id=m.id)} for m in recent],
        'urls': {'contacts_missing': url_for('comms.contacts', missing=1),
                 'compose': url_for('comms.compose'),
                 'compose_sms': url_for('comms.compose', channel='SMS'),
                 'compose_email': url_for('comms.compose', channel='Email'),
                 'announcements': url_for('comms.announcements'),
                 'templates': url_for('comms.templates_list'),
                 'messages': url_for('comms.messages_list')},
    })


# ============================================================================
# CONTACT DIRECTORY
# ============================================================================

@comms_bp.route('/announcements')
@login_required
def announcements():
    items = Announcement.query.order_by(Announcement.is_pinned.desc(),
                                        Announcement.created_at.desc()).all()
    # Acknowledgement counts (only meaningful for needs_ack announcements).
    from models import AnnouncementAck
    ack_ids = [a.id for a in items if a.needs_ack]
    ack_counts = {}
    if ack_ids:
        ack_counts = dict(db.session.query(AnnouncementAck.announcement_id,
                                           func.count(AnnouncementAck.id))
                          .filter(AnnouncementAck.announcement_id.in_(ack_ids))
                          .group_by(AnnouncementAck.announcement_id).all())
    return _render({
        'page': 'announcements', 'nav': _nav_urls(),
        'items': [{'id': a.id, 'title': a.title, 'body': a.body or '',
                   'category': a.category or 'Info', 'audience': a.audience or 'All',
                   'is_pinned': bool(a.is_pinned), 'is_active': bool(a.is_active),
                   'needs_ack': bool(a.needs_ack), 'ack_count': ack_counts.get(a.id, 0),
                   'attachment': _attachment_dict(a.attachment_id),
                   'created_at': a.created_at.strftime('%d %b %Y') if a.created_at else '',
                   'created_by': a.created_by or '',
                   'starts_on': a.starts_on.strftime('%d %b') if a.starts_on else '',
                   'ends_on': a.ends_on.strftime('%d %b') if a.ends_on else '',
                   'delete_url': url_for('comms.delete_announcement', ann_id=a.id)} for a in items],
        'add_url': url_for('comms.add_announcement'),
        'upload_url': url_for('comms.upload_attachment'),
    })


@comms_bp.route('/announcements/<int:ann_id>/ack', methods=['POST'])
@login_required
def ack_announcement(ann_id):
    """Record that the current user has acknowledged an announcement (idempotent)."""
    from models import AnnouncementAck
    from utils.access_control import get_current_user
    from sqlalchemy.exc import IntegrityError
    user = get_current_user()
    if user is None:
        return _err('Sign in to acknowledge.', url_for('main.dashboard'))
    if not AnnouncementAck.query.filter_by(announcement_id=ann_id, user_id=user.id).first():
        db.session.add(AnnouncementAck(announcement_id=ann_id, user_id=user.id))
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()          # raced a duplicate — already acknowledged
    return _ok('Acknowledged.', url_for('main.dashboard'))


def _read_announcement(a):
    a.title = strip_tags(request.form.get('title'))
    a.body = strip_tags(request.form.get('body')) or None
    a.audience = request.form.get('audience') or 'All'
    a.category = request.form.get('category') or 'Info'
    a.is_pinned = bool(request.form.get('is_pinned'))
    a.needs_ack = bool(request.form.get('needs_ack'))
    a.starts_on = _date(request.form.get('starts_on'))
    a.ends_on = _date(request.form.get('ends_on'))
    att_id = request.form.get('attachment_id', type=int)
    if att_id is not None:               # empty = clear, a value = set
        a.attachment_id = att_id or None


def _date(v):
    try:
        return datetime.strptime(v, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


@comms_bp.route('/announcements/add', methods=['POST'])
@login_required
def add_announcement():
    if not (request.form.get('title') or '').strip():
        return _err('Title is required.', url_for('comms.announcements'))
    a = Announcement(created_by=_current_user())
    _read_announcement(a)
    db.session.add(a)
    db.session.commit()
    return _ok('Announcement posted.', url_for('comms.announcements'))


@comms_bp.route('/announcements/<int:ann_id>/edit', methods=['POST'])
@login_required
def edit_announcement(ann_id):
    a = db.get_or_404(Announcement, ann_id)
    _read_announcement(a)
    db.session.commit()
    return _ok('Announcement updated.', url_for('comms.announcements'))


@comms_bp.route('/announcements/<int:ann_id>/delete', methods=['POST'])
@login_required
def delete_announcement(ann_id):
    a = db.get_or_404(Announcement, ann_id)
    from utils.audit import log_action
    log_action('communication.announcement_delete', target=a)
    db.session.delete(a)
    db.session.commit()
    return _ok('Announcement deleted.', url_for('comms.announcements'))


@comms_bp.route('/contacts')
@login_required
def contacts():
    term = _term_from(request.args.get('term_id', type=int))
    class_id = request.args.get('class_id', type=int)
    q = (request.args.get('q') or '').strip()
    missing = request.args.get('missing')  # only students without a contact

    terms = Term.query.order_by(Term.id.desc()).all()
    classes = SchoolClass.query.filter_by(is_active=True).order_by(SchoolClass.level).all()

    if class_id and term:
        ids = [e.student_id for e in (StudentEnrollment.query
               .join(ClassArmAssignment,
                     StudentEnrollment.class_arm_assignment_id == ClassArmAssignment.id)
               .filter(StudentEnrollment.is_active == True,
                       ClassArmAssignment.term_id == term.id,
                       ClassArmAssignment.class_id == class_id).all())]
        base = Student.query.filter(Student.id.in_(ids or [-1]))
    else:
        base = Student.query.filter_by(is_active=True)
    from utils.branch_scope import scope_query
    from utils.access_control import teacher_form_student_ids
    base = scope_query(base, Student)
    form_ids = teacher_form_student_ids()
    if form_ids is not None:
        base = base.filter(Student.id.in_(form_ids or [-1]))
    if q:
        like = like_term(q)
        base = base.filter(db.or_(Student.surname.ilike(like, escape='\\'),
                                  Student.first_name.ilike(like, escape='\\'),
                                  Student.student_id.ilike(like, escape='\\')))
    students = base.order_by(Student.surname, Student.first_name).limit(400).all()

    rows = []
    for s in students:
        contacts_list = [c for c in s.parent_contacts.all() if c.phone_number]
        if missing and contacts_list:
            continue
        rows.append({
            'student': {'id': s.id, 'full_name': s.full_name, 'student_id': s.student_id,
                        'view_url': url_for('main.view_student', student_id=s.id)},
            'contacts': [{'name': c.name or 'Parent', 'relationship': c.relationship or '',
                          'is_primary': bool(c.is_primary), 'phone_number': c.phone_number,
                          'wa_intl': comms.normalise_phone(c.phone_number)} for c in contacts_list],
        })

    return _render({
        'page': 'contacts', 'nav': _nav_urls(), 'rows': rows,
        'terms': [{'id': t.id, 'full_name': t.full_name} for t in terms],
        'classes': [{'id': c.id, 'name': c.name} for c in classes],
        'term_id': term.id if term else '', 'class_id': class_id or '',
        'q': q, 'missing': bool(missing), 'cov': comms.coverage_stats(),
        'self_url': url_for('comms.contacts'),
    })


# ============================================================================
# TEMPLATES
# ============================================================================

@comms_bp.route('/templates')
@login_required
def templates_list():
    q = (request.args.get('q') or '').strip()
    category = (request.args.get('category') or '').strip()
    base = MessageTemplate.query
    if q:
        like = like_term(q)
        base = base.filter(db.or_(MessageTemplate.name.ilike(like, escape='\\'),
                                  MessageTemplate.body.ilike(like, escape='\\'),
                                  MessageTemplate.category.ilike(like, escape='\\')))
    if category:
        base = base.filter(MessageTemplate.category == category)
    # Favourites first, then active, then by category / name.
    tpls = base.order_by(MessageTemplate.is_favorite.desc(),
                         MessageTemplate.is_active.desc(),
                         MessageTemplate.category, MessageTemplate.name).all()
    categories = sorted({c[0] for c in MessageTemplate.query
                         .with_entities(MessageTemplate.category).distinct().all() if c[0]})
    return _render({
        'page': 'templates', 'nav': _nav_urls(), 'placeholders': comms.PLACEHOLDERS,
        'categories': categories,
        'sel': {'q': q, 'category': category},
        'templates': [{'id': t.id, 'name': t.name, 'body': t.body,
                       'category': t.category or '', 'is_active': bool(t.is_active),
                       'is_favorite': bool(t.is_favorite),
                       'use_url': url_for('comms.compose', tpl=t.id),
                       'edit_url': url_for('comms.edit_template', template_id=t.id),
                       'duplicate_url': url_for('comms.duplicate_template', template_id=t.id),
                       'favorite_url': url_for('comms.toggle_favorite', template_id=t.id),
                       'delete_url': url_for('comms.delete_template', template_id=t.id)} for t in tpls],
        'add_url': url_for('comms.add_template'),
        'self_url': url_for('comms.templates_list'),
    })


@comms_bp.route('/templates/add', methods=['POST'])
@login_required
def add_template():
    name = (request.form.get('name') or '').strip()
    body = (request.form.get('body') or '').strip()
    if not (name and body):
        return _err('Template name and body are required.', url_for('comms.templates_list'))
    db.session.add(MessageTemplate(name=name, body=body,
        category=(request.form.get('category') or 'General').strip()))
    db.session.commit()
    return _ok(f'Template "{name}" saved.', url_for('comms.templates_list'))


@comms_bp.route('/templates/<int:template_id>/edit', methods=['POST'])
@login_required
def edit_template(template_id):
    t = db.get_or_404(MessageTemplate, template_id)
    t.name = (request.form.get('name') or t.name).strip()
    t.body = (request.form.get('body') or t.body).strip()
    t.category = (request.form.get('category') or t.category or 'General').strip()
    t.is_active = bool(request.form.get('is_active'))
    db.session.commit()
    return _ok('Template updated.', url_for('comms.templates_list'))


@comms_bp.route('/templates/<int:template_id>/delete', methods=['POST'])
@login_required
def delete_template(template_id):
    t = db.get_or_404(MessageTemplate, template_id)
    db.session.delete(t)
    db.session.commit()
    return _ok('Template deleted.', url_for('comms.templates_list'))


@comms_bp.route('/templates/<int:template_id>/duplicate', methods=['POST'])
@login_required
def duplicate_template(template_id):
    t = db.get_or_404(MessageTemplate, template_id)
    copy = MessageTemplate(name=f'{t.name} (copy)', body=t.body, category=t.category,
                           is_active=t.is_active, is_favorite=False)
    db.session.add(copy)
    db.session.commit()
    return _ok(f'Duplicated "{t.name}".', url_for('comms.templates_list'))


@comms_bp.route('/templates/<int:template_id>/favorite', methods=['POST'])
@login_required
def toggle_favorite(template_id):
    t = db.get_or_404(MessageTemplate, template_id)
    t.is_favorite = not bool(t.is_favorite)
    db.session.commit()
    return _ok('Favourite updated.' if t.is_favorite else 'Removed from favourites.',
               url_for('comms.templates_list'))


# ============================================================================
# INBOX — internal staff-to-staff messaging
# ============================================================================

def _me():
    from utils.access_control import get_current_user
    return get_current_user()


@comms_bp.route('/inbox')
@login_required
def inbox():
    """The staff messaging inbox — conversation list (the thread loads client-side)."""
    from utils import chat
    me = _me()
    if me is None:
        return _err('Messaging is for staff accounts.', url_for('comms.dashboard'))
    return _render({
        'page': 'inbox', 'nav': _nav_urls(), 'me': {'id': me.id, 'name': me.full_name or me.username},
        'conversations': chat.conversations_for(me.id),
        'urls': {'self': url_for('comms.inbox'), 'start': url_for('comms.inbox_start'),
                 'users': url_for('comms.inbox_users'), 'upload': url_for('comms.upload_attachment'),
                 'thread': url_for('comms.inbox_thread', conv_id=0)[:-1]},
    })


@comms_bp.route('/inbox/users')
@login_required
def inbox_users():
    """Search staff users to start a conversation with (branch-scoped, never self)."""
    from models import User
    from utils.branch_scope import scope_query, is_central, viewing_branch_id
    me = _me()
    if me is None:
        return jsonify([])
    q = (request.args.get('q') or '').strip()
    base = User.query.filter(User.is_active.is_(True), User.id != me.id)
    if not is_central():
        bid = viewing_branch_id()
        if bid not in (None, -1):
            base = base.filter(User.branch_id == bid)
    if q:
        like = like_term(q)
        base = base.filter(db.or_(User.full_name.ilike(like, escape='\\'),
                                  User.username.ilike(like, escape='\\')))
    rows = base.order_by(User.full_name, User.username).limit(20).all()
    return jsonify([{'id': u.id, 'name': u.full_name or u.username,
                     'label': f'{u.full_name or u.username} ({u.role})'} for u in rows])


@comms_bp.route('/inbox/start', methods=['POST'])
@login_required
def inbox_start():
    """Start (or reuse) a direct chat, or create a group. Returns the thread url."""
    from utils import chat
    from utils.branch_scope import branch_for_new
    me = _me()
    if me is None:
        return _err('Messaging is for staff accounts.', url_for('comms.dashboard'))
    ids = request.form.getlist('user_ids', type=int)
    if not ids:
        return _err('Pick at least one person.', url_for('comms.inbox'))
    if len(ids) == 1:
        conv = chat.get_or_create_direct(me.id, ids[0], branch_id=branch_for_new())
    else:
        conv = chat.create_group(me.id, ids, request.form.get('title'), branch_id=branch_for_new())
    if not conv:
        return _err('Could not start that conversation.', url_for('comms.inbox'))
    return _ok('Conversation ready.', url_for('comms.inbox_thread', conv_id=conv.id))


@comms_bp.route('/inbox/<int:conv_id>')
@login_required
def inbox_thread(conv_id):
    """A conversation's messages (marks it read for the viewer)."""
    from utils import chat
    from models import Conversation, ChatMessage
    me = _me()
    if me is None or not chat.is_member(conv_id, me.id):
        return _err('You are not part of that conversation.', url_for('comms.inbox'))
    conv = db.get_or_404(Conversation, conv_id)
    msgs = ChatMessage.query.filter_by(conversation_id=conv_id).order_by(ChatMessage.created_at).all()
    chat.mark_read(conv_id, me.id)
    rows = [{'id': m.id, 'body': m.body or '', 'mine': m.sender_id == me.id,
             'sender': chat._display_name(m.sender) if m.sender else '—',
             'at': m.created_at.strftime('%d %b %H:%M') if m.created_at else '',
             'attachment': _attachment_dict(m.attachment_id)} for m in msgs]
    return _render({
        'page': 'inbox', 'nav': _nav_urls(), 'me': {'id': me.id, 'name': me.full_name or me.username},
        'conversations': chat.conversations_for(me.id),
        'active': {'id': conv.id, 'title': chat.conversation_title(conv, me.id),
                   'kind': conv.kind, 'messages': rows,
                   'send_url': url_for('comms.inbox_send', conv_id=conv.id)},
        'urls': {'self': url_for('comms.inbox'), 'start': url_for('comms.inbox_start'),
                 'users': url_for('comms.inbox_users'), 'upload': url_for('comms.upload_attachment'),
                 'thread': url_for('comms.inbox_thread', conv_id=0)[:-1]},
    })


@comms_bp.route('/inbox/<int:conv_id>/send', methods=['POST'])
@login_required
def inbox_send(conv_id):
    from utils import chat
    me = _me()
    if me is None or not chat.is_member(conv_id, me.id):
        return _err('You are not part of that conversation.', url_for('comms.inbox'))
    m = chat.post_message(conv_id, me.id, request.form.get('body'),
                          attachment_id=request.form.get('attachment_id', type=int) or None)
    if not m:
        return _err('Message cannot be empty.', url_for('comms.inbox_thread', conv_id=conv_id))
    return _ok('Sent.', url_for('comms.inbox_thread', conv_id=conv_id))


@comms_bp.route('/inbox/unread-count')
@login_required
def inbox_unread():
    from utils import chat
    me = _me()
    return jsonify({'unread': chat.total_unread(me.id) if me else 0})


# ============================================================================
# ATTACHMENTS (announcements + email campaigns)
# ============================================================================

@comms_bp.route('/attachments', methods=['POST'])
@login_required
def upload_attachment():
    """Store an uploaded file and return its metadata for the composer / announcement
    form to reference by id."""
    from utils import comm_attachments as CA
    try:
        att = CA.save(request.files.get('file'), created_by=_current_user())
    except ValueError as e:
        return _err(str(e), url_for('comms.compose'))
    from utils.audit import log_action
    log_action('communication.attachment_upload', target=att, detail=att.original_name)
    return jsonify({'ok': True, 'attachment': CA.as_dict(
        att, download_url=url_for('comms.download_attachment', att_id=att.id))})


@comms_bp.route('/attachments/<int:att_id>')
@login_required
def download_attachment(att_id):
    """Stream an attachment from the tenant's upload folder."""
    from flask import send_file
    from utils import comm_attachments as CA
    from models import CommAttachment
    att = db.get_or_404(CommAttachment, att_id)
    path = CA.fs_path(att)
    if not path:
        return ('File not found.', 404)
    return send_file(path, as_attachment=True, download_name=att.original_name,
                     mimetype=att.content_type or 'application/octet-stream')


def _attachment_dict(att_id):
    """Serialise a referenced attachment (or None) for a payload."""
    if not att_id:
        return None
    from utils import comm_attachments as CA
    from models import CommAttachment
    att = db.session.get(CommAttachment, att_id)
    return CA.as_dict(att, download_url=url_for('comms.download_attachment', att_id=att.id)) if att else None


# ============================================================================
# COMPOSE / SEND
# ============================================================================

@comms_bp.route('/compose', methods=['GET', 'POST'])
@login_required
def compose():
    term = _term_from(request.values.get('term_id', type=int))
    terms = Term.query.order_by(Term.id.desc()).all()
    classes = SchoolClass.query.filter_by(is_active=True).order_by(SchoolClass.level).all()
    arms = ClassArm.query.filter_by(is_active=True, is_default=False).order_by(ClassArm.name).all()
    templates = MessageTemplate.query.filter_by(is_active=True).order_by(MessageTemplate.name).all()

    if request.method == 'POST':
        body = strip_tags(request.form.get('body'))
        channel = request.form.get('channel') or 'SMS'
        title = strip_tags(request.form.get('title'))
        spec = _recipient_spec(request.form)

        if not body:
            return _err('Message body cannot be empty.', url_for('comms.compose'))
        if spec['to'] == 'staff' and not _is_admin():
            return _err('Only administrators can message staff.', url_for('comms.compose'))

        # How many are reachable on this channel (phone for SMS/WhatsApp, email for
        # Email) — decides the count in the label and whether there's anyone to send.
        targets = comms.resolve_recipients(spec, term)
        reachable = comms.reachable_targets(targets, channel)
        if not reachable:
            miss = 'an email address' if comms.channel_is_email(channel) else 'a phone number'
            who = 'staff' if spec['to'] == 'staff' else 'recipients'
            return _err(f'No {who} with {miss} matched that selection.',
                        url_for('comms.compose'))

        # Optional scheduling (auto-send via the matching gateway at a future time).
        from utils import sms_gateway, mailer
        gateway_ready = mailer.is_configured() if comms.channel_is_email(channel) \
            else sms_gateway.is_configured()
        scheduled_at, status = None, 'Draft'
        if request.form.get('schedule') and gateway_ready:
            scheduled_at = _dt(request.form.get('scheduled_at'))
            if scheduled_at and scheduled_at > datetime.now():
                status = 'Scheduled'

        from utils.branch_scope import branch_for_new
        label = _spec_label(spec, classes, arms, len(reachable))
        att_id = request.form.get('attachment_id', type=int) or None
        msg = comms.build_campaign(
            body, channel=channel, spec=spec, term=term, title=title,
            audience_label=label, created_by=_current_user(),
            status=status, scheduled_at=scheduled_at, branch_id=branch_for_new(),
            attachment_id=att_id)
        if comms.channel_is_inapp(channel):
            note = f'In-app notification sent to {len(reachable)} staff member(s).'
        elif status == 'Scheduled':
            note = (f'Campaign scheduled for {scheduled_at.strftime("%d %b %Y, %I:%M %p")} '
                    f'({len(reachable)} recipient(s)).')
        else:
            note = f'Campaign created for {len(reachable)} recipient(s).'
        return _ok(note, url_for('comms.message_detail', message_id=msg.id))

    # Pre-selection from query params (e.g. "Message defaulters" from Finance,
    # or "Use" from the templates page).
    pre_audience = request.args.get('audience') or 'all'
    pre_class = request.args.get('class_id', type=int)
    pre_tpl = request.args.get('tpl', type=int)
    pre_body = ''
    if pre_tpl:
        t = db.session.get(MessageTemplate, pre_tpl)
        if t:
            pre_body = t.body
    elif pre_audience == 'defaulters':
        # Convenience: default to the Fee Reminder template when messaging debtors.
        t = MessageTemplate.query.filter(MessageTemplate.name.ilike('%fee%reminder%'),
                                         MessageTemplate.is_active == True).first()
        if t:
            pre_tpl, pre_body = t.id, t.body
    elif request.args.get('notice') == 'results':
        # Pre-select the result-notification template when releasing results.
        t = MessageTemplate.query.filter(MessageTemplate.name.ilike('%result%'),
                                         MessageTemplate.is_active == True).first()
        if t:
            pre_tpl, pre_body = t.id, t.body

    from utils import sms_gateway, mailer
    from models import Department
    gw = sms_gateway.get_config()
    departments = (Department.query.filter_by(is_active=True).order_by(Department.name).all()
                   if _is_admin() else [])
    return _render({
        'page': 'compose', 'nav': _nav_urls(), 'is_admin': _is_admin(),
        'term_id': term.id if term else '',
        'terms': [{'id': t.id, 'full_name': t.full_name} for t in terms],
        'classes': [{'id': c.id, 'name': c.name} for c in classes],
        'arms': [{'id': a.id, 'name': a.name} for a in arms],
        'departments': [{'id': d.id, 'name': d.name} for d in departments],
        'streams': ['Science', 'Arts', 'Commercial'],
        'genders': ['Male', 'Female'],
        'templates': [{'id': t.id, 'name': t.name, 'category': t.category or '', 'body': t.body}
                      for t in templates],
        # In-app posts to staff bells, so it's offered only where staff messaging is.
        'channels': (CHANNELS if _is_admin() else [c for c in CHANNELS if c != 'In-app']),
        'placeholders': comms.PLACEHOLDERS, 'cov': comms.coverage_stats(),
        'gateway_ready': sms_gateway.is_configured(gw),
        'gateway_label': sms_gateway.provider_label(gw),
        'email_ready': mailer.is_configured(),
        'pre_channel': (request.args.get('channel') if request.args.get('channel') in CHANNELS else ''),
        'pre_audience': pre_audience, 'pre_class': pre_class or '',
        'pre_tpl': pre_tpl or '', 'pre_body': pre_body,
        'groups': _saved_groups(),
        'urls': {'submit': url_for('comms.compose'), 'preview': url_for('comms.compose_preview'),
                 'search': url_for('comms.students_search'), 'settings': url_for('comms.settings'),
                 'save_group': url_for('comms.save_group'), 'upload': url_for('comms.upload_attachment')},
    })


def _current_user():
    from flask import session
    return session.get('username') or session.get('user') or 'Admin'


def _recipient_spec(form):
    """Parse the composer's recipient controls into a spec dict for
    comms.resolve_recipients (parents-or-staff + filters + exclusions)."""
    to = (form.get('to') or 'parents').lower()
    if to == 'staff':
        return {
            'to': 'staff',
            'staff_scope': (form.get('staff_scope') or 'all'),
            'department_id': form.get('department_id', type=int),
            'exclude_ids': form.getlist('exclude_ids', type=int),
        }
    return {
        'to': 'parents',
        'audience': form.get('audience') or 'all',
        'class_id': form.get('class_id', type=int),
        'arm_id': form.get('arm_id', type=int),
        'student_ids': form.getlist('student_ids', type=int),
        'gender': (form.get('gender') or '').strip(),
        'stream': (form.get('stream') or '').strip(),
        'exclude_ids': form.getlist('exclude_ids', type=int),
    }


def _spec_label(spec, classes, arms, count):
    """Human label for a recipient spec (shown as the campaign's audience label)."""
    if spec.get('to') == 'staff':
        scope = (spec.get('staff_scope') or 'all')
        base = {'teaching': 'Teaching staff', 'non-teaching': 'Non-teaching staff',
                'department': 'Department staff'}.get(scope, 'All staff')
        return f'{base} ({count})'
    audience = spec.get('audience') or 'all'
    suffix = []
    if spec.get('gender'):
        suffix.append(spec['gender'])
    if spec.get('stream'):
        suffix.append(spec['stream'])
    extra = (' · ' + ', '.join(suffix)) if suffix else ''
    if audience == 'all':
        return f'All parents{extra} ({count})'
    if audience == 'defaulters':
        base = 'Fee defaulters'
        cls = next((c for c in classes if c.id == spec.get('class_id')), None)
        if cls:
            base += f' · {cls.name}'
        return f'{base}{extra} ({count})'
    if audience in ('class', 'arm'):
        cls = next((c for c in classes if c.id == spec.get('class_id')), None)
        arm = next((a for a in arms if a.id == spec.get('arm_id')), None)
        name = ' '.join(p for p in [cls.name if cls else '', arm.name if arm else ''] if p)
        return f'{name or "Class"}{extra} ({count})'
    return f'Selected students{extra} ({count})'


@comms_bp.route('/compose/preview', methods=['POST'])
@login_required
def compose_preview():
    """JSON: recipient counts (channel-aware) + a personalised sample for the
    chosen recipient selection."""
    term = _term_from(request.form.get('term_id', type=int))
    channel = request.form.get('channel') or 'SMS'
    body = request.form.get('body') or ''
    spec = _recipient_spec(request.form)
    if spec['to'] == 'staff' and not _is_admin():
        return jsonify({'total': 0, 'reachable': 0, 'no_phone': 0, 'unreachable': 0,
                        'by_email': comms.channel_is_email(channel), 'sample': '',
                        'denied': True})
    targets = comms.resolve_recipients(spec, term)
    reachable = comms.reachable_targets(targets, channel)
    sample = ''
    if reachable:
        sample = comms.render(body, comms.campaign_context(reachable[0], term))
    return jsonify({'total': len(targets), 'reachable': len(reachable),
                    'no_phone': len(targets) - len(reachable),
                    'unreachable': len(targets) - len(reachable),
                    'by_email': comms.channel_is_email(channel), 'sample': sample})


@comms_bp.route('/students/search')
@login_required
def students_search():
    q = (request.args.get('q') or '').strip()
    if len(q) < 2:
        return jsonify([])
    like = like_term(q)
    from utils.branch_scope import scope_query
    rows = (scope_query(Student.query.filter_by(is_active=True), Student)
            .filter(db.or_(Student.surname.ilike(like, escape='\\'), Student.first_name.ilike(like, escape='\\'),
                           Student.student_id.ilike(like, escape='\\')))
            .order_by(Student.surname).limit(15).all())
    return jsonify([{'id': s.id, 'name': s.full_name, 'sid': s.student_id,
                     'label': f'{s.full_name} ({s.student_id})'} for s in rows])


# ============================================================================
# SAVED RECIPIENT GROUPS
# ============================================================================

# The composer only persists these spec keys — never free-form data — so a saved
# group can't smuggle anything unexpected back into resolve_recipients.
_SPEC_KEYS = ('to', 'audience', 'class_id', 'arm_id', 'student_ids', 'gender',
              'stream', 'staff_scope', 'department_id', 'exclude_ids')


def _clean_spec(form):
    spec = _recipient_spec(form)
    return {k: v for k, v in spec.items() if k in _SPEC_KEYS and v not in (None, '', [])}


def _saved_groups():
    from models import RecipientGroup
    from utils.branch_scope import scope_query
    rows = (scope_query(RecipientGroup.query, RecipientGroup)
            .order_by(RecipientGroup.name).all())
    return [{'id': g.id, 'name': g.name, 'spec': g.spec_dict(),
             'delete_url': url_for('comms.delete_group', group_id=g.id)} for g in rows]


@comms_bp.route('/groups/save', methods=['POST'])
@login_required
def save_group():
    import json
    from models import RecipientGroup
    from utils.branch_scope import branch_for_new
    name = (request.form.get('name') or '').strip()
    if not name:
        return _err('Give the group a name.', url_for('comms.compose'))
    spec = _clean_spec(request.form)
    if spec.get('to') == 'staff' and not _is_admin():
        return _err('Only administrators can save staff groups.', url_for('comms.compose'))
    g = (RecipientGroup.query.filter_by(name=name).first()
         if request.form.get('overwrite') else None)
    if g is None:
        g = RecipientGroup(name=name, branch_id=branch_for_new(), created_by=_current_user())
        db.session.add(g)
    g.spec = json.dumps(spec)
    db.session.commit()
    from utils.audit import log_action
    log_action('communication.group_saved', target=g, detail=name)
    return _ok(f'Saved recipient group "{name}".', url_for('comms.compose'))


@comms_bp.route('/groups/<int:group_id>/delete', methods=['POST'])
@login_required
def delete_group(group_id):
    from models import RecipientGroup
    g = db.get_or_404(RecipientGroup, group_id)
    require_branch_access(g.branch_id)
    db.session.delete(g)
    db.session.commit()
    return _ok('Recipient group deleted.', url_for('comms.compose'))


@comms_bp.route('/messages/<int:message_id>/cancel-schedule', methods=['POST'])
@login_required
def cancel_schedule(message_id):
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    msg.status = 'Draft'
    msg.scheduled_at = None
    db.session.commit()
    return _ok('Schedule cancelled — the campaign is now on hold.',
               url_for('comms.message_detail', message_id=message_id))


@comms_bp.route('/process-scheduled', methods=['POST'])
@admin_required
def process_scheduled():
    n = comms.dispatch_due_scheduled()
    return _ok(f'Processed {n} due scheduled campaign(s).' if n else 'No campaigns were due.',
               url_for('comms.messages_list'))


# ============================================================================
# COMMUNICATION HISTORY — one searchable timeline (campaigns + announcements)
# ============================================================================

HISTORY_TYPES = ['SMS', 'WhatsApp', 'Email', 'Announcement']
HISTORY_STATUSES = ['Draft', 'Scheduled', 'Sending', 'Sent']
_HISTORY_PAGE_SIZE = 25
_HISTORY_CAP = 2000            # per-source scan bound (schools are modest volume)


def _history_filters():
    a = request.args
    return {
        'type': a.get('type') or '',
        'status': a.get('status') or '',
        'sender': a.get('sender') or '',
        'q': (a.get('q') or '').strip(),
        'from': _date(a.get('from')),
        'to': _date(a.get('to')),
        'page': max(a.get('page', type=int) or 1, 1),
    }


def _history_ids(f):
    """Lightweight (kind, id, created_at) rows for both sources under the filters —
    merged + sorted so we only hydrate the current page."""
    end = None
    if f['to']:
        end = datetime.combine(f['to'], datetime.max.time())
    start = datetime.combine(f['from'], datetime.min.time()) if f['from'] else None
    rows = []

    include_campaigns = (not f['type']) or f['type'] in ('SMS', 'WhatsApp', 'Email')
    if include_campaigns:
        mq = scope_query(db.session.query(Message.id, Message.created_at), Message)
        if f['type'] in ('SMS', 'WhatsApp', 'Email'):
            mq = mq.filter(Message.channel == f['type'])
        if f['status']:
            mq = mq.filter(Message.status == f['status'])
        if f['sender']:
            mq = mq.filter(Message.created_by == f['sender'])
        if start:
            mq = mq.filter(Message.created_at >= start)
        if end:
            mq = mq.filter(Message.created_at <= end)
        if f['q']:
            like = like_term(f['q'])
            mq = mq.filter(db.or_(Message.title.ilike(like, escape='\\'),
                                  Message.body.ilike(like, escape='\\'),
                                  Message.audience_label.ilike(like, escape='\\')))
        rows += [('campaign', i, c) for i, c in
                 mq.order_by(Message.created_at.desc()).limit(_HISTORY_CAP).all()]

    # Announcements are school-wide (no branch/status/channel); include them only
    # when the type filter allows and no campaign-only status is selected.
    include_ann = (f['type'] in ('', 'Announcement')) and not f['status']
    if include_ann:
        aq = db.session.query(Announcement.id, Announcement.created_at)
        if f['sender']:
            aq = aq.filter(Announcement.created_by == f['sender'])
        if start:
            aq = aq.filter(Announcement.created_at >= start)
        if end:
            aq = aq.filter(Announcement.created_at <= end)
        if f['q']:
            like = like_term(f['q'])
            aq = aq.filter(db.or_(Announcement.title.ilike(like, escape='\\'),
                                  Announcement.body.ilike(like, escape='\\')))
        rows += [('announcement', i, c) for i, c in
                 aq.order_by(Announcement.created_at.desc()).limit(_HISTORY_CAP).all()]

    rows.sort(key=lambda r: (r[2] or datetime.min), reverse=True)
    return rows


def _history_item_campaign(m):
    return {'kind': 'campaign', 'id': m.id, 'type': m.channel,
            'date': m.created_at.strftime('%d %b %Y') if m.created_at else '',
            'title': m.title or 'Message', 'status': m.status,
            'scheduled_at': m.scheduled_at.strftime('%d %b %H:%M') if m.scheduled_at else '',
            'audience_label': m.audience_label or '', 'by': m.created_by or '',
            'recipient_count': m.recipient_count or 0, 'sent_count': m.sent_count or 0,
            'url': url_for('comms.message_detail', message_id=m.id)}


def _history_item_announcement(a):
    return {'kind': 'announcement', 'id': a.id, 'type': 'Announcement',
            'date': a.created_at.strftime('%d %b %Y') if a.created_at else '',
            'title': a.title or 'Announcement', 'status': 'Posted',
            'scheduled_at': '', 'audience_label': a.audience or 'All',
            'by': a.created_by or '', 'recipient_count': '', 'sent_count': '',
            'url': url_for('comms.announcements')}


@comms_bp.route('/messages')
@login_required
def messages_list():
    f = _history_filters()
    ids = _history_ids(f)
    total = len(ids)
    pages = max((total + _HISTORY_PAGE_SIZE - 1) // _HISTORY_PAGE_SIZE, 1)
    page = min(f['page'], pages)
    window = ids[(page - 1) * _HISTORY_PAGE_SIZE: page * _HISTORY_PAGE_SIZE]

    camp_ids = [i for k, i, _ in window if k == 'campaign']
    ann_ids = [i for k, i, _ in window if k == 'announcement']
    camps = {m.id: m for m in Message.query.filter(Message.id.in_(camp_ids)).all()} if camp_ids else {}
    anns = {a.id: a for a in Announcement.query.filter(Announcement.id.in_(ann_ids)).all()} if ann_ids else {}
    items = []
    for kind, i, _c in window:
        if kind == 'campaign' and i in camps:
            items.append(_history_item_campaign(camps[i]))
        elif kind == 'announcement' and i in anns:
            items.append(_history_item_announcement(anns[i]))

    # Sender options (distinct authors across both sources, branch-scoped campaigns).
    senders = set(x[0] for x in scope_query(
        db.session.query(Message.created_by).filter(Message.created_by.isnot(None)), Message).distinct().all())
    senders |= set(x[0] for x in db.session.query(Announcement.created_by)
                   .filter(Announcement.created_by.isnot(None)).distinct().all())

    return _render({
        'page': 'messages', 'nav': _nav_urls(), 'is_admin': _is_admin(),
        'items': items, 'total': total, 'page_no': page, 'pages': pages,
        'has_prev': page > 1, 'has_next': page < pages,
        'types': HISTORY_TYPES, 'statuses': HISTORY_STATUSES,
        'senders': sorted(s for s in senders if s),
        'sel': {'type': f['type'], 'status': f['status'], 'sender': f['sender'],
                'q': f['q'], 'from': request.args.get('from', ''), 'to': request.args.get('to', ''),
                'page': page},
        'urls': {'compose': url_for('comms.compose'), 'self': url_for('comms.messages_list'),
                 'reports': url_for('comms.reports'),
                 'process_scheduled': url_for('comms.process_scheduled')},
    })


# ============================================================================
# REPORTS — usage & delivery analytics over a date range
# ============================================================================

def _report_range():
    """(from_date, to_date) for reports — defaults to the last 30 days."""
    import datetime as _dt
    to = _date(request.args.get('to')) or _dt.date.today()
    frm = _date(request.args.get('from')) or (to - _dt.timedelta(days=30))
    return frm, to


def _report_data(frm, to):
    """Usage + delivery aggregates for campaigns created in [frm, to]."""
    start = datetime.combine(frm, datetime.min.time())
    end = datetime.combine(to, datetime.max.time())
    mq = scope_query(Message.query, Message).filter(
        Message.created_at >= start, Message.created_at <= end)

    by_channel = {}
    total_campaigns = 0
    for m in mq.all():
        total_campaigns += 1
        c = by_channel.setdefault(m.channel or 'Other',
                                  {'channel': m.channel or 'Other', 'campaigns': 0,
                                   'recipients': 0, 'sent': 0})
        c['campaigns'] += 1
        c['recipients'] += m.recipient_count or 0
        c['sent'] += m.sent_count or 0
    channel_rows = sorted(by_channel.values(), key=lambda r: -r['campaigns'])

    # Recipient-level delivery, scoped by the owning campaign + date window.
    rq = (scope_query(db.session.query(MessageRecipient.status, func.count(MessageRecipient.id))
                      .join(Message, MessageRecipient.message_id == Message.id), Message)
          .filter(Message.created_at >= start, Message.created_at <= end)
          .group_by(MessageRecipient.status))
    by_status = {s: n for s, n in rq.all()}
    sent = by_status.get('Sent', 0)
    failed = by_status.get('Failed', 0)
    pending = by_status.get('Pending', 0)
    attempted = sent + failed
    delivery_rate = round(sent / attempted * 100, 1) if attempted else None
    # Read receipts (currently in-app; email opens land here in future).
    read = (scope_query(db.session.query(func.count(MessageRecipient.id))
                        .join(Message, MessageRecipient.message_id == Message.id), Message)
            .filter(Message.created_at >= start, Message.created_at <= end,
                    MessageRecipient.read_at.isnot(None)).scalar() or 0)
    read_rate = round(read / sent * 100, 1) if sent else None

    scheduled = mq.filter(Message.status == 'Scheduled').count()
    drafts = mq.filter(Message.status == 'Draft').count()
    announcements = (db.session.query(func.count(Announcement.id))
                     .filter(Announcement.created_at >= start,
                             Announcement.created_at <= end).scalar() or 0)
    return {
        'from': frm.strftime('%Y-%m-%d'), 'to': to.strftime('%Y-%m-%d'),
        'total_campaigns': total_campaigns, 'by_channel': channel_rows,
        'recipients': sum(c['recipients'] for c in channel_rows),
        'sent': sent, 'failed': failed, 'pending': pending,
        'delivery_rate': delivery_rate, 'scheduled': scheduled, 'drafts': drafts,
        'read': int(read), 'read_rate': read_rate,
        'announcements': int(announcements),
    }


@comms_bp.route('/reports')
@login_required
def reports():
    frm, to = _report_range()
    data = _report_data(frm, to)
    return _render({
        'page': 'reports', 'nav': _nav_urls(), 'is_admin': _is_admin(),
        'data': data, 'sel': {'from': data['from'], 'to': data['to']},
        'urls': {'self': url_for('comms.reports'),
                 'export_csv': url_for('comms.reports_export', format='csv'),
                 'export_xlsx': url_for('comms.reports_export', format='xlsx'),
                 'history': url_for('comms.messages_list')},
    })


@comms_bp.route('/reports/export')
@login_required
def reports_export():
    frm, to = _report_range()
    d = _report_data(frm, to)
    from utils.audit import log_action
    log_action('data.export_comm_report', detail=f'{d["from"]}..{d["to"]}')
    headers = ['Channel', 'Campaigns', 'Recipients', 'Sent']
    summary = [
        ['Report period', f'{d["from"]} to {d["to"]}'],
        ['Total campaigns', d['total_campaigns']],
        ['Total recipients', d['recipients']],
        ['Delivered (sent)', d['sent']],
        ['Failed', d['failed']],
        ['Pending', d['pending']],
        ['Delivery rate', ('%s%%' % d['delivery_rate']) if d['delivery_rate'] is not None else 'n/a'],
        ['Scheduled', d['scheduled']],
        ['Drafts', d['drafts']],
        ['Announcements', d['announcements']],
    ]
    if request.args.get('format') == 'xlsx':
        from openpyxl import Workbook
        from utils.web_exports import xlsx_response
        wb = Workbook()
        ws = wb.active; ws.title = 'Summary'
        for row in summary:
            ws.append(row)
        ws2 = wb.create_sheet('By channel')
        ws2.append(headers)
        for c in d['by_channel']:
            ws2.append([c['channel'], c['campaigns'], c['recipients'], c['sent']])
        return xlsx_response(wb, f'comm_report_{d["from"]}_{d["to"]}.xlsx')
    out = io.StringIO()
    w = csv.writer(out)
    for row in summary:
        w.writerow(row)
    w.writerow([])
    w.writerow(headers)
    for c in d['by_channel']:
        w.writerow([c['channel'], c['campaigns'], c['recipients'], c['sent']])
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition':
                             f'attachment; filename=comm_report_{d["from"]}_{d["to"]}.csv'})


@comms_bp.route('/messages/<int:message_id>')
@login_required
def message_detail(message_id):
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    recips = msg.recipients.order_by(MessageRecipient.parent_name).all()
    is_email = (msg.channel or '').lower() == 'email'
    rows = []
    for r in recips:
        # For email campaigns the destination is the address; reuse the existing
        # "phone" column so the React table shows it without a bundle change.
        dest = (r.email or '') if is_email else r.phone
        rows.append({'id': r.id, 'parent_name': r.parent_name,
                     'student_name': r.student.full_name if r.student else '—',
                     'phone': dest, 'email': r.email or '',
                     'intl': '' if is_email else comms.normalise_phone(r.phone),
                     'status': r.status, 'error': r.error or '', 'body': r.body,
                     'read': r.read_at is not None,
                     'sent_url': url_for('comms.mark_sent', message_id=msg.id, rid=r.id)})
    from utils import sms_gateway, mailer
    gw = sms_gateway.get_config()
    is_inapp = comms.channel_is_inapp(msg.channel)
    if is_inapp:                       # bell notifications are delivered on creation
        channel_ready, channel_label = False, 'In-app'
    elif is_email:
        channel_ready, channel_label = mailer.is_configured(), 'Email'
    else:
        channel_ready, channel_label = sms_gateway.is_configured(gw), sms_gateway.provider_label(gw)
    return _render({
        'page': 'message_detail', 'nav': _nav_urls(), 'is_admin': _is_admin(),
        'msg': {'id': msg.id, 'title': msg.title or 'Campaign', 'channel': msg.channel,
                'audience_label': msg.audience_label, 'status': msg.status, 'body': msg.body,
                'scheduled_at': msg.scheduled_at.strftime('%d %b, %I:%M %p') if msg.scheduled_at else '',
                'created_at': msg.created_at.strftime('%d %b %Y, %I:%M %p') if msg.created_at else '',
                'created_by': msg.created_by or '', 'sent_count': msg.sent_count or 0,
                'recipient_count': msg.recipient_count or 0,
                'attachment': _attachment_dict(msg.attachment_id)},
        'rows': rows, 'segments': comms.sms_segments(msg.body),
        'gateway_ready': channel_ready,
        'gateway_label': channel_label,
        'failed_count': msg.recipients.filter(MessageRecipient.status == 'Failed').count(),
        'pending_count': msg.recipients.filter(MessageRecipient.status != 'Sent').count(),
        'read_count': msg.recipients.filter(MessageRecipient.read_at.isnot(None)).count(),
        'urls': {'export': url_for('comms.export_recipients', message_id=msg.id),
                 'compose': url_for('comms.compose'),
                 'cancel_schedule': url_for('comms.cancel_schedule', message_id=msg.id),
                 'send_gateway': url_for('comms.send_gateway', message_id=msg.id),
                 'mark_all_sent': url_for('comms.mark_all_sent', message_id=msg.id),
                 'delete': url_for('comms.delete_message', message_id=msg.id),
                 'list': url_for('comms.messages_list')},
    })


@comms_bp.route('/messages/<int:message_id>/recipient/<int:rid>/sent', methods=['POST'])
@login_required
def mark_sent(message_id, rid):
    r = db.get_or_404(MessageRecipient, rid)
    if r.message_id != message_id:
        return ('', 404)
    require_branch_access(r.message.branch_id)
    if r.status != 'Sent':
        r.status = 'Sent'
        r.sent_at = datetime.now()
        r.message.sent_count = (r.message.sent_count or 0) + 1
        db.session.commit()
    if request.headers.get('X-Requested-With') == 'fetch':
        return jsonify({'ok': True, 'sent': r.message.sent_count})
    return redirect(url_for('comms.message_detail', message_id=message_id))


@comms_bp.route('/messages/<int:message_id>/mark-all-sent', methods=['POST'])
@login_required
def mark_all_sent(message_id):
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    n = 0
    for r in msg.recipients.filter(MessageRecipient.status != 'Sent').all():
        r.status = 'Sent'
        r.sent_at = datetime.now()
        n += 1
    msg.sent_count = msg.recipients.filter_by(status='Sent').count()
    db.session.commit()
    return _ok(f'Marked {n} recipient(s) as sent.',
               url_for('comms.message_detail', message_id=message_id))


@comms_bp.route('/messages/<int:message_id>/export')
@login_required
def export_recipients(message_id):
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    from utils.audit import log_action
    log_action('data.export_recipients', target=msg,
               detail='campaign recipients (parent names/phones)')
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(['Parent', 'Phone', 'Phone (intl)', 'Student', 'Message', 'Status'])
    from utils.web_exports import formula_guard as _fg
    for r in msg.recipients.all():
        w.writerow([_fg(r.parent_name), _fg(r.phone), comms.normalise_phone(r.phone),
                    _fg(r.student.full_name if r.student else ''), _fg(r.body), r.status])
    fname = f'campaign_{msg.id}_recipients.csv'
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': f'attachment; filename={fname}'})


@comms_bp.route('/messages/<int:message_id>/delete', methods=['POST'])
@admin_required
def delete_message(message_id):
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    from utils.audit import log_action
    log_action('communication.message_delete',
               target_type='message', target_id=msg.id, target_label=getattr(msg, 'title', None))
    db.session.delete(msg)
    db.session.commit()
    return _ok('Campaign deleted.', url_for('comms.messages_list'))


@comms_bp.route('/messages/<int:message_id>/send-gateway', methods=['POST'])
@login_required
def send_gateway(message_id):
    """Dispatch all pending recipients through the configured gateway — the SMS
    provider for SMS/WhatsApp campaigns, or SMTP for Email campaigns."""
    from utils import sms_gateway, mailer
    from flask import current_app
    msg = db.get_or_404(Message, message_id)
    require_branch_access(msg.branch_id)
    if comms.channel_is_inapp(msg.channel):
        return _err('In-app notifications are delivered instantly — nothing to send.',
                    url_for('comms.message_detail', message_id=message_id))
    is_email = (msg.channel or '').lower() == 'email'

    if is_email:
        if not mailer.is_configured():
            return _err('Email is not configured. Set your SMTP details before sending email.',
                        url_for('comms.message_detail', message_id=message_id))
    else:
        cfg = sms_gateway.get_config()
        if not sms_gateway.is_configured(cfg):
            return _err('No SMS gateway is configured. Add your provider key in Settings.',
                        url_for('comms.message_detail', message_id=message_id))

    # Claim a scheduled campaign so this manual send can't race the worker.
    if msg.status in ('Scheduled', 'Sending') and not comms._claim_message(msg.id, msg.status):
        return _err('This campaign is already being sent.',
                    url_for('comms.message_detail', message_id=message_id))

    # Send in the background: each gateway/SMTP call can take seconds, so sending a
    # whole batch inline would tie up this web worker for minutes.
    pending = msg.recipients.filter(MessageRecipient.status != 'Sent').count()
    msg.status = 'Sending'
    db.session.commit()
    if is_email:
        comms.dispatch_campaign_email_async(current_app._get_current_object(), msg.id)
        via = 'email'
    else:
        comms.dispatch_campaign_async(current_app._get_current_object(), msg.id, cfg)
        via = sms_gateway.provider_label(cfg)
    return _ok(f'Sending {pending} message(s) via {via} in the '
               'background. Refresh this page to see delivery progress.',
               url_for('comms.message_detail', message_id=message_id))


# ============================================================================
# SMS GATEWAY SETTINGS
# ============================================================================

@comms_bp.route('/settings')
@login_required
def settings():
    from utils import sms_gateway, automations
    cfg = sms_gateway.get_config()
    configured = sms_gateway.is_configured(cfg)
    balance_ok, balance = (sms_gateway.get_balance(cfg) if configured else (False, ''))
    return _render({
        'page': 'settings', 'nav': _nav_urls(), 'is_admin': _is_admin(),
        'cfg': {'provider': cfg.get('provider', 'none'), 'sender': cfg.get('sender', ''),
                'termii_key': cfg.get('termii_key', ''), 'twilio_sid': cfg.get('twilio_sid', ''),
                'twilio_token': cfg.get('twilio_token', '')},
        'providers': [{'key': k, 'label': v} for k, v in sms_gateway.PROVIDERS.items()],
        'configured': configured, 'balance_ok': balance_ok, 'balance': balance,
        'provider_label': sms_gateway.provider_label(cfg),
        'automations': automations.all_states(),
        'urls': {'save': url_for('comms.save_settings'), 'test': url_for('comms.test_sms'),
                 'save_automations': url_for('comms.save_automations')},
    })


@comms_bp.route('/settings/automations', methods=['POST'])
@admin_required
def save_automations():
    """Persist the per-automation on/off toggles (unchecked box = disabled)."""
    from utils import automations
    for key in automations.KEYS:
        automations.set_enabled(key, request.form.get(key) == 'on')
    return _ok('Automation settings saved.', url_for('comms.settings'))


@comms_bp.route('/settings/save', methods=['POST'])
@admin_required
def save_settings():
    from utils import sms_gateway
    sms_gateway.save_config(request.form)
    cfg = sms_gateway.get_config()
    if cfg['provider'] != 'none' and not sms_gateway.is_configured(cfg):
        return _ok('Settings saved, but some required fields for this provider are missing.',
                   url_for('comms.settings'))
    return _ok('SMS gateway settings saved.', url_for('comms.settings'))


@comms_bp.route('/settings/test', methods=['POST'])
@admin_required
def test_sms():
    from utils import sms_gateway
    phone = (request.form.get('phone') or '').strip()
    if not phone:
        return _err('Enter a phone number to send the test to.', url_for('comms.settings'))
    ok, info = sms_gateway.send_sms(phone,
        f'Test message from {comms.school_name()} via PosyHub. Your SMS gateway is working!')
    if ok:
        return _ok(f'Test SMS sent successfully (ref: {info}).', url_for('comms.settings'))
    return _err(f'Test failed: {info}', url_for('comms.settings'))
